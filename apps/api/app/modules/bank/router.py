import csv
import io
from collections.abc import Iterable
from datetime import datetime
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import get_session
from app.models import (
    BankTransaction,
    ExpenseBankMatch,
    Ledger,
    LedgerStatus,
    RevenueBankMatch,
    SyncJob,
    SyncJobStatus,
    User,
    utc_now,
)
from app.modules.audit.service import write_audit_log
from app.modules.auth.permissions import (
    ensure_permission,
    ensure_store_access,
    scoped_store_condition,
)
from app.modules.auth.router import audit_actor, get_current_user
from app.modules.common import paginate
from app.schemas import (
    ApiEnvelope,
    BankImportPreviewResult,
    BankImportPreviewRow,
    BankImportResult,
    BankImportRollbackResult,
    BankImportRowError,
    BankTransactionCreate,
    BankTransactionRead,
    BankTransactionUpdate,
    Page,
)

router = APIRouter(prefix="/bank-transactions", tags=["bank"])

BANK_IMPORT_TEMPLATE_HEADERS = ["发生时间", "类型", "金额", "备注", "流水号"]
BANK_IMPORT_TEMPLATE_ROWS = [
    ["2026-08-20 10:00:00", "收入", "1200.00", "营业款", "BANK-EXAMPLE-001"],
    ["2026-08-21 11:30:00", "支出", "300.00", "物料款", "BANK-EXAMPLE-002"],
]


def ensure_open_ledger(session: Session, store_id: str, period: str) -> None:
    ledger = session.scalar(select(Ledger).where(Ledger.store_id == store_id, Ledger.period == period))
    if ledger is None:
        raise HTTPException(status_code=404, detail="Ledger not found")
    if ledger.status == LedgerStatus.CLOSED.value:
        raise HTTPException(status_code=409, detail="Ledger is closed")


def ensure_open_or_create_ledger(session: Session, store_id: str, period: str) -> None:
    ledger = session.scalar(select(Ledger).where(Ledger.store_id == store_id, Ledger.period == period))
    if ledger is None:
        session.add(Ledger(store_id=store_id, period=period))
        session.flush()
        return
    if ledger.status == LedgerStatus.CLOSED.value:
        raise HTTPException(status_code=409, detail="Ledger is closed")


def normalize_bank_assignment(session: Session, payload: BankTransactionCreate) -> dict:
    data = payload.model_dump()
    if data["store_id"] and not data["ledger_period"]:
        data["ledger_period"] = data["occurred_at"].strftime("%Y-%m")
        ensure_open_or_create_ledger(session, data["store_id"], data["ledger_period"])
    elif data["store_id"] and data["ledger_period"]:
        ensure_open_or_create_ledger(session, data["store_id"], data["ledger_period"])
    elif data["ledger_period"]:
        raise HTTPException(status_code=422, detail="Store is required when ledger period is provided")
    return data


@router.get("", response_model=ApiEnvelope[Page[BankTransactionRead]])
def list_bank_transactions(
    store_id: str | None = None,
    ledger_period: str | None = None,
    direction: str | None = None,
    page: int = 1,
    page_size: int = 50,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ApiEnvelope[Page[BankTransactionRead]]:
    ensure_permission(session, current_user, "reconciliation.view")
    query = select(BankTransaction).order_by(BankTransaction.occurred_at.desc())
    if store_id:
        ensure_store_access(session, current_user, store_id)
        query = query.where(BankTransaction.store_id == store_id)
    else:
        query = query.where(scoped_store_condition(session, current_user, BankTransaction.store_id))
    if ledger_period:
        query = query.where(BankTransaction.ledger_period == ledger_period)
    if direction:
        query = query.where(BankTransaction.direction == direction)
    items, total = paginate(session, query, page, page_size)
    return ApiEnvelope(data=Page(items=items, total=total, page=page, page_size=page_size))


@router.post("", response_model=ApiEnvelope[BankTransactionRead], status_code=201)
def create_bank_transaction(
    payload: BankTransactionCreate,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ApiEnvelope[BankTransactionRead]:
    ensure_permission(session, current_user, "reconciliation.manage")
    if not payload.store_id:
        raise HTTPException(status_code=422, detail="Store is required")
    ensure_store_access(session, current_user, payload.store_id)
    transaction = BankTransaction(**normalize_bank_assignment(session, payload))
    session.add(transaction)
    session.flush()
    write_audit_log(
        session,
        actor=audit_actor(current_user),
        action="bank_transaction.create",
        resource_type="bank_transaction",
        resource_id=transaction.id,
        summary=f"新增银行流水：{transaction.amount}",
        metadata={
            "store_id": transaction.store_id,
            "ledger_period": transaction.ledger_period,
            "direction": transaction.direction,
        },
    )
    session.commit()
    session.refresh(transaction)
    return ApiEnvelope(data=transaction)


@router.get("/import/template.csv")
def download_bank_import_template(
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    ensure_permission(session, current_user, "reconciliation.manage")
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(BANK_IMPORT_TEMPLATE_HEADERS)
    writer.writerows(BANK_IMPORT_TEMPLATE_ROWS)
    content = "\ufeff" + output.getvalue()
    filename = "bank-import-template.csv"
    return StreamingResponse(
        iter([content.encode("utf-8")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.patch("/{transaction_id}", response_model=ApiEnvelope[BankTransactionRead])
def update_bank_transaction(
    transaction_id: str,
    payload: BankTransactionUpdate,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ApiEnvelope[BankTransactionRead]:
    ensure_permission(session, current_user, "reconciliation.manage")
    transaction = session.get(BankTransaction, transaction_id)
    if transaction is None:
        raise HTTPException(status_code=404, detail="Bank transaction not found")
    ensure_store_access(session, current_user, transaction.store_id)

    updates = payload.model_dump(exclude_unset=True)
    target_store_id = updates.get("store_id", transaction.store_id)
    ensure_store_access(session, current_user, target_store_id)
    target_ledger_period = updates.get("ledger_period", transaction.ledger_period)
    if target_store_id and not target_ledger_period:
        target_ledger_period = updates.get("occurred_at", transaction.occurred_at).strftime("%Y-%m")
        updates["ledger_period"] = target_ledger_period
        ensure_open_or_create_ledger(session, target_store_id, target_ledger_period)
    elif target_store_id and target_ledger_period:
        ensure_open_or_create_ledger(session, target_store_id, target_ledger_period)
    elif target_ledger_period:
        raise HTTPException(status_code=422, detail="Store is required when ledger period is provided")
    new_amount = updates.get("amount")
    if new_amount is not None and Decimal(new_amount) < Decimal(transaction.matched_amount or 0):
        raise HTTPException(status_code=409, detail="Amount cannot be lower than matched amount")

    new_serial_no = updates.get("bank_serial_no")
    if new_serial_no:
        exists = session.scalar(
            select(BankTransaction).where(
                BankTransaction.bank_serial_no == new_serial_no,
                BankTransaction.id != transaction_id,
            )
        )
        if exists is not None:
            raise HTTPException(status_code=409, detail="Bank serial number already exists")

    for field, value in updates.items():
        setattr(transaction, field, value)

    write_audit_log(
        session,
        actor=audit_actor(current_user),
        action="bank_transaction.update",
        resource_type="bank_transaction",
        resource_id=transaction.id,
        summary=f"更新银行流水：{transaction.amount}",
        metadata={
            "store_id": transaction.store_id,
            "ledger_period": transaction.ledger_period,
            "updated_fields": sorted(updates.keys()),
        },
    )
    session.commit()
    session.refresh(transaction)
    return ApiEnvelope(data=transaction)


@router.delete("/{transaction_id}", response_model=ApiEnvelope[BankTransactionRead])
def delete_bank_transaction(
    transaction_id: str,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ApiEnvelope[BankTransactionRead]:
    ensure_permission(session, current_user, "reconciliation.manage")
    transaction = session.get(BankTransaction, transaction_id)
    if transaction is None:
        raise HTTPException(status_code=404, detail="Bank transaction not found")
    ensure_store_access(session, current_user, transaction.store_id)
    if Decimal(transaction.matched_amount or 0) > 0:
        raise HTTPException(status_code=409, detail="Bank transaction already matched")
    has_expense_match = session.scalar(
        select(ExpenseBankMatch).where(ExpenseBankMatch.bank_transaction_id == transaction_id).limit(1)
    )
    has_revenue_match = session.scalar(
        select(RevenueBankMatch).where(RevenueBankMatch.bank_transaction_id == transaction_id).limit(1)
    )
    if has_expense_match or has_revenue_match:
        raise HTTPException(status_code=409, detail="Bank transaction already has match records")

    deleted = BankTransactionRead.model_validate(transaction)
    write_audit_log(
        session,
        actor=audit_actor(current_user),
        action="bank_transaction.delete",
        resource_type="bank_transaction",
        resource_id=transaction.id,
        summary=f"删除银行流水：{transaction.amount}",
        metadata={
            "store_id": transaction.store_id,
            "ledger_period": transaction.ledger_period,
            "direction": transaction.direction,
        },
    )
    session.delete(transaction)
    session.commit()
    return ApiEnvelope(data=deleted)


def pick(row: dict[str, str | None], *names: str) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and value.strip():
            return value.strip()
    return ""


def parse_direction(value: str) -> str:
    normalized = value.strip().lower()
    if normalized in {"income", "in", "收入", "入账", "收款"}:
        return "income"
    if normalized in {"expense", "out", "支出", "出账", "付款"}:
        return "expense"
    raise ValueError(f"Unsupported direction: {value}")


def parse_datetime(value: str) -> datetime:
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, pattern)
        except ValueError:
            continue
    return datetime.fromisoformat(value)


def transaction_exists(session: Session, payload: dict) -> bool:
    serial_no = payload.get("bank_serial_no")
    if serial_no:
        exists = session.scalar(select(BankTransaction).where(BankTransaction.bank_serial_no == serial_no))
        return exists is not None
    exists = session.scalar(
        select(BankTransaction).where(
            BankTransaction.store_id == payload.get("store_id"),
            BankTransaction.ledger_period == payload.get("ledger_period"),
            BankTransaction.occurred_at == payload["occurred_at"],
            BankTransaction.direction == payload["direction"],
            BankTransaction.amount == payload["amount"],
            BankTransaction.counterparty_name == payload.get("counterparty_name"),
        )
    )
    return exists is not None


def parse_import_payload(row: dict[str, str | None], store_id: str | None, ledger_period: str | None) -> dict:
    amount_text = pick(row, "amount", "金额", "交易金额")
    occurred_at = parse_datetime(pick(row, "occurred_at", "发生时间", "交易时间", "日期"))
    return {
        "store_id": store_id,
        "ledger_period": ledger_period or (occurred_at.strftime("%Y-%m") if store_id else None),
        "occurred_at": occurred_at,
        "direction": parse_direction(pick(row, "direction", "方向", "类型", "收支方向", "收入还是支出")),
        "amount": Decimal(amount_text.replace(",", "")),
        "counterparty_name": pick(row, "counterparty_name", "对方户名", "交易对方") or None,
        "counterparty_account": pick(row, "counterparty_account", "对方账号") or None,
        "summary": pick(row, "summary", "摘要", "备注") or None,
        "bank_serial_no": pick(row, "bank_serial_no", "流水号", "交易流水号") or None,
    }


def read_import_rows(file_name: str | None, content: bytes) -> Iterable[tuple[int, dict[str, str | None]]]:
    lower_name = (file_name or "").lower()
    if lower_name.endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        if not any(headers):
            raise HTTPException(status_code=400, detail="XLSX header is required")
        for index, values in enumerate(rows, start=2):
            yield index, {
                headers[column_index]: str(value).strip() if value is not None else None
                for column_index, value in enumerate(values)
                if column_index < len(headers) and headers[column_index]
            }
        workbook.close()
        return

    if lower_name.endswith(".csv") or not lower_name:
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise HTTPException(status_code=400, detail="CSV header is required")
        for index, row in enumerate(reader, start=2):
            yield index, row
        return

    raise HTTPException(status_code=400, detail="Only CSV and XLSX files are supported")


async def preview_bank_transactions_file(
    store_id: str | None = Form(None),
    ledger_period: str | None = Form(None),
    file: UploadFile = File(...),
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ApiEnvelope[BankImportPreviewResult]:
    ensure_permission(session, current_user, "reconciliation.manage")
    ensure_store_access(session, current_user, store_id)
    if ledger_period and not store_id:
        raise HTTPException(status_code=422, detail="Store is required when ledger period is provided")
    if store_id and ledger_period:
        ensure_open_or_create_ledger(session, store_id, ledger_period)
    content = await file.read()

    duplicate_count = 0
    row_errors: list[BankImportRowError] = []
    preview_rows: list[BankImportPreviewRow] = []
    for index, row in read_import_rows(file.filename, content):
        try:
            payload = parse_import_payload(row, store_id, ledger_period)
            duplicate = transaction_exists(session, payload)
            if duplicate:
                duplicate_count += 1
            preview_rows.append(
                BankImportPreviewRow(
                    row_number=index,
                    occurred_at=payload["occurred_at"],
                    direction=payload["direction"],
                    amount=payload["amount"],
                    counterparty_name=payload["counterparty_name"],
                    counterparty_account=payload["counterparty_account"],
                    summary=payload["summary"],
                    bank_serial_no=payload["bank_serial_no"],
                    duplicate=duplicate,
                )
            )
        except (ValueError, InvalidOperation) as exc:
            row_errors.append(BankImportRowError(row_number=index, message=str(exc)))

    return ApiEnvelope(
        data=BankImportPreviewResult(
            valid_count=len(preview_rows) - duplicate_count,
            duplicate_count=duplicate_count,
            error_count=len(row_errors),
            preview_rows=preview_rows[:20],
            row_errors=row_errors,
        )
    )


async def import_bank_transactions_file(
    store_id: str | None = Form(None),
    ledger_period: str | None = Form(None),
    started_by: str = Form("admin"),
    file: UploadFile = File(...),
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ApiEnvelope[BankImportResult]:
    ensure_permission(session, current_user, "reconciliation.manage")
    ensure_store_access(session, current_user, store_id)
    if ledger_period and not store_id:
        raise HTTPException(status_code=422, detail="Store is required when ledger period is provided")
    if store_id and ledger_period:
        ensure_open_or_create_ledger(session, store_id, ledger_period)
    content = await file.read()

    job = SyncJob(
        job_type="bank_transaction_import",
        status=SyncJobStatus.RUNNING.value,
        started_by=audit_actor(current_user, started_by),
        started_at=utc_now(),
    )
    session.add(job)
    session.flush()

    created_count = 0
    skipped_count = 0
    row_errors: list[BankImportRowError] = []
    try:
        for index, row in read_import_rows(file.filename, content):
            job.processed_count += 1
            try:
                payload = parse_import_payload(row, store_id, ledger_period)
                if payload["store_id"] and payload["ledger_period"] and not ledger_period:
                    ensure_open_or_create_ledger(session, payload["store_id"], payload["ledger_period"])
                if transaction_exists(session, payload):
                    skipped_count += 1
                    continue
                payload["import_job_id"] = job.id
                session.add(BankTransaction(**payload))
                created_count += 1
            except (ValueError, InvalidOperation) as exc:
                row_errors.append(BankImportRowError(row_number=index, message=str(exc)))
                job.failed_count += 1
                job.error_message = "; ".join(
                    f"Row {error.row_number}: {error.message}" for error in row_errors[:5]
                )
        job.success_count = created_count
        job.failed_count += skipped_count
        job.status = SyncJobStatus.SUCCEEDED.value if not job.error_message else SyncJobStatus.FAILED.value
        job.finished_at = utc_now()
        write_audit_log(
            session,
            actor=audit_actor(current_user, started_by),
            action="bank_transaction.import_csv",
            resource_type="sync_job",
            resource_id=job.id,
            summary=f"导入银行流水：新增 {created_count} 条，跳过 {skipped_count} 条",
            metadata={"store_id": store_id, "ledger_period": ledger_period, "file_name": file.filename},
        )
        session.commit()
        session.refresh(job)
    except Exception:
        job.status = SyncJobStatus.FAILED.value
        job.finished_at = utc_now()
        session.commit()
        raise

    return ApiEnvelope(
        data=BankImportResult(
            job=job,
            created_count=created_count,
            skipped_count=skipped_count,
            row_errors=row_errors,
        )
    )


router.post("/import/preview", response_model=ApiEnvelope[BankImportPreviewResult])(
    preview_bank_transactions_file
)
router.post("/import", response_model=ApiEnvelope[BankImportResult], status_code=201)(
    import_bank_transactions_file
)
router.post("/import-csv", response_model=ApiEnvelope[BankImportResult], status_code=201)(
    import_bank_transactions_file
)


@router.post("/imports/{job_id}/rollback", response_model=ApiEnvelope[BankImportRollbackResult])
def rollback_bank_import(
    job_id: str,
    operator: str = "admin",
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ApiEnvelope[BankImportRollbackResult]:
    ensure_permission(session, current_user, "reconciliation.manage")
    job = session.get(SyncJob, job_id)
    if job is None or job.job_type != "bank_transaction_import":
        raise HTTPException(status_code=404, detail="Bank import job not found")

    transactions = list(
        session.scalars(select(BankTransaction).where(BankTransaction.import_job_id == job_id))
    )
    if not transactions:
        return ApiEnvelope(data=BankImportRollbackResult(job=job, deleted_count=0))
    store_ids = {transaction.store_id for transaction in transactions if transaction.store_id}
    for store_id in store_ids:
        ensure_store_access(session, current_user, store_id)
    transaction_ids = [transaction.id for transaction in transactions]
    matched_transactions = [transaction for transaction in transactions if Decimal(transaction.matched_amount or 0) > 0]
    if matched_transactions:
        raise HTTPException(status_code=409, detail="Imported transactions already matched")
    has_expense_match = session.scalar(
        select(ExpenseBankMatch).where(ExpenseBankMatch.bank_transaction_id.in_(transaction_ids)).limit(1)
    )
    has_revenue_match = session.scalar(
        select(RevenueBankMatch).where(RevenueBankMatch.bank_transaction_id.in_(transaction_ids)).limit(1)
    )
    if has_expense_match or has_revenue_match:
        raise HTTPException(status_code=409, detail="Imported transactions already have match records")

    deleted_count = len(transactions)
    for transaction in transactions:
        session.delete(transaction)
    job.status = SyncJobStatus.FAILED.value
    job.error_message = "Rolled back by operator"
    job.finished_at = utc_now()
    write_audit_log(
        session,
        actor=audit_actor(current_user, operator),
        action="bank_transaction_import.rollback",
        resource_type="sync_job",
        resource_id=job.id,
        summary=f"回滚银行流水导入：删除 {deleted_count} 条",
        metadata={"deleted_count": deleted_count},
    )
    session.commit()
    session.refresh(job)
    return ApiEnvelope(data=BankImportRollbackResult(job=job, deleted_count=deleted_count))
