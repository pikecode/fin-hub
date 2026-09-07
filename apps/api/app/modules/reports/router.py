import csv
import io
from datetime import datetime
from decimal import Decimal

from fastapi import APIRouter, Depends, Header, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.core.database import get_session
from app.models import (
    ApprovalInstance,
    ApprovalTemplate,
    BankTransaction,
    ExpenseBankMatch,
    ExpenseItem,
    Ledger,
    MatchStatus,
    RevenueBankMatch,
    RevenueBankMatchRecord,
    RevenueRecord,
    ShareholderAccessGrant,
    Store,
    User,
)
from app.modules.approvals.status import approval_expense_stats_map
from app.modules.auth.permissions import effective_store_ids, ensure_permission, ensure_store_access
from app.modules.auth.router import get_optional_current_user
from app.modules.shareholder_auth.service import grant_store_ids, require_shareholder_grant
from app.schemas import (
    ApiEnvelope,
    FinancialAnalyticsApprovalStatusItem,
    ExpenseBreakdownItem,
    FinancialAnalyticsCategoryItem,
    FinancialAnalyticsDetailReport,
    FinancialAnalyticsMetrics,
    FinancialAnalyticsReconciliationItem,
    FinancialAnalyticsReport,
    FinancialAnalyticsStoreItem,
    FinancialAnalyticsTemplateItem,
    FinancialAnalyticsTrendItem,
    RevenueChannelMonthlyBreakdownItem,
    LedgerPeriodOption,
    LedgerReportDetail,
    LedgerReportSummary,
    LedgerTrend,
    ReconciliationRecord,
    ReportPeriodOption,
    RevenueChannelBreakdownItem,
    StoreComparisonReport,
    StoreReportSummary,
)

router = APIRouter(prefix="/reports", tags=["reports"])


def read_bearer_token(authorization: str | None) -> str | None:
    if not authorization:
        return None
    prefix = "Bearer "
    if not authorization.startswith(prefix):
        raise HTTPException(status_code=401, detail="Invalid authorization header")
    return authorization.removeprefix(prefix)


def get_report_access(
    authorization: str | None = Header(default=None),
    current_user: User | None = Depends(get_optional_current_user),
    session: Session = Depends(get_session),
) -> ShareholderAccessGrant | None:
    if current_user is not None:
        ensure_permission(session, current_user, "reports.view")
        return None
    token = read_bearer_token(authorization)
    if token is None:
        raise HTTPException(status_code=401, detail="Not authenticated")
    return require_shareholder_grant(session, token)


def ensure_report_store_access(
    session: Session,
    grant: ShareholderAccessGrant | None,
    current_user: User | None,
    store_id: str,
) -> None:
    if grant is None:
        if current_user is not None:
            ensure_store_access(session, current_user, store_id)
        return
    if store_id not in grant_store_ids(grant):
        raise HTTPException(status_code=403, detail="Store is not authorized")


def filter_report_stores(
    session: Session,
    stores: list[Store],
    shareholder_grant: ShareholderAccessGrant | None,
    current_user: User | None,
) -> list[Store]:
    if shareholder_grant is not None:
        allowed_store_ids = set(grant_store_ids(shareholder_grant))
        return [store for store in stores if store.id in allowed_store_ids]
    if current_user is None:
        return stores
    allowed_store_ids = set(effective_store_ids(session, current_user))
    return [store for store in stores if store.id in allowed_store_ids]


def ensure_shareholder_can_read_ledger(grant: ShareholderAccessGrant | None, ledger: Ledger) -> None:
    if grant is not None and ledger.status != "closed":
        raise HTTPException(status_code=403, detail="Ledger is not closed")


def decimal_sum(value: Decimal | None) -> Decimal:
    return value if value is not None else Decimal("0.00")


def decimal_cell(value: Decimal | None) -> float:
    return float(decimal_sum(value))


def decimal_rate(numerator: Decimal, denominator: Decimal) -> Decimal:
    if denominator <= 0:
        return Decimal("0.00")
    return (numerator / denominator * Decimal(100)).quantize(Decimal("0.01"))


def build_report_summary(session: Session, ledger: Ledger, store: Store) -> LedgerReportSummary:
    revenue_amount = decimal_sum(
        session.scalar(
            select(func.sum(RevenueRecord.gross_amount)).where(
                RevenueRecord.store_id == ledger.store_id,
                RevenueRecord.ledger_period == ledger.period,
            )
        )
    )
    bank_income_amount = decimal_sum(
        session.scalar(
            select(func.sum(BankTransaction.amount)).where(
                BankTransaction.store_id == ledger.store_id,
                BankTransaction.ledger_period == ledger.period,
                BankTransaction.direction == "income",
            )
        )
    )
    income_amount = revenue_amount if revenue_amount > 0 else bank_income_amount
    expense_amount = decimal_sum(
        session.scalar(
            select(func.sum(ExpenseItem.amount)).where(
                ExpenseItem.store_id == ledger.store_id,
                ExpenseItem.ledger_period == ledger.period,
            )
        )
    )
    pending_expense_count = session.scalar(
        select(func.count()).select_from(ExpenseItem).where(
            ExpenseItem.store_id == ledger.store_id,
            ExpenseItem.ledger_period == ledger.period,
            ExpenseItem.payment_status.in_(["unpaid", "partial_paid"]),
        )
    )
    pending_bank_count = session.scalar(
        select(func.count()).select_from(BankTransaction).where(
            BankTransaction.store_id == ledger.store_id,
            BankTransaction.ledger_period == ledger.period,
            BankTransaction.matched_amount < BankTransaction.amount,
        )
    )
    return LedgerReportSummary(
        store_id=store.id,
        store_name=store.name,
        period=ledger.period,
        ledger_status=ledger.status,
        income_amount=income_amount,
        expense_amount=expense_amount,
        profit_amount=income_amount - expense_amount,
        pending_expense_count=pending_expense_count or 0,
        pending_bank_transaction_count=pending_bank_count or 0,
    )


def read_ledger_for_report(
    session: Session,
    store_id: str,
    period: str,
    shareholder_grant: ShareholderAccessGrant | None,
    current_user: User | None = None,
) -> tuple[Store, Ledger]:
    ensure_report_store_access(session, shareholder_grant, current_user, store_id)
    store = session.get(Store, store_id)
    if store is None:
        raise HTTPException(status_code=404, detail="Store not found")
    ledger = session.scalar(select(Ledger).where(Ledger.store_id == store_id, Ledger.period == period))
    if ledger is None:
        raise HTTPException(status_code=404, detail="Ledger not found")
    ensure_shareholder_can_read_ledger(shareholder_grant, ledger)
    return store, ledger


def read_ledger_export_rows(
    session: Session,
    store_id: str,
    period: str,
) -> tuple[list[RevenueRecord], list[ExpenseItem], list[BankTransaction]]:
    revenue_records = session.scalars(
        select(RevenueRecord)
        .where(RevenueRecord.store_id == store_id, RevenueRecord.ledger_period == period)
        .order_by(RevenueRecord.revenue_date.asc(), RevenueRecord.created_at.asc())
    ).all()
    expenses = session.scalars(
        select(ExpenseItem)
        .where(ExpenseItem.store_id == store_id, ExpenseItem.ledger_period == period)
        .order_by(ExpenseItem.expense_date.asc(), ExpenseItem.created_at.asc())
    ).all()
    bank_transactions = session.scalars(
        select(BankTransaction)
        .where(BankTransaction.store_id == store_id, BankTransaction.ledger_period == period)
        .order_by(BankTransaction.occurred_at.asc())
    ).all()
    return list(revenue_records), list(expenses), list(bank_transactions)


def build_revenue_channel_breakdown(
    revenue_records: list[RevenueRecord],
    revenue_matches: list[RevenueBankMatch],
) -> list[RevenueChannelBreakdownItem]:
    buckets: dict[str, dict[str, Decimal | int]] = {}
    for record in revenue_records:
        bucket = buckets.setdefault(
            record.channel,
            {
                "gross_amount": Decimal("0.00"),
                "net_amount": Decimal("0.00"),
                "fee_amount": Decimal("0.00"),
                "matched_amount": Decimal("0.00"),
                "record_count": 0,
            },
        )
        bucket["gross_amount"] = Decimal(bucket["gross_amount"]) + Decimal(record.gross_amount)
        bucket["net_amount"] = Decimal(bucket["net_amount"]) + Decimal(record.net_amount)
        bucket["fee_amount"] = Decimal(bucket["fee_amount"]) + Decimal(record.fee_amount)
        bucket["record_count"] = int(bucket["record_count"]) + 1

    for match in revenue_matches:
        if match.status != MatchStatus.CONFIRMED.value:
            continue
        bucket = buckets.setdefault(
            match.channel,
            {
                "gross_amount": Decimal("0.00"),
                "net_amount": Decimal("0.00"),
                "fee_amount": Decimal("0.00"),
                "matched_amount": Decimal("0.00"),
                "record_count": 0,
            },
        )
        bucket["matched_amount"] = Decimal(bucket["matched_amount"]) + Decimal(match.amount)

    breakdown = [
        RevenueChannelBreakdownItem(
            channel=channel,
            gross_amount=Decimal(bucket["gross_amount"]),
            net_amount=Decimal(bucket["net_amount"]),
            fee_amount=Decimal(bucket["fee_amount"]),
            fee_rate=decimal_rate(Decimal(bucket["fee_amount"]), Decimal(bucket["gross_amount"])),
            matched_amount=Decimal(bucket["matched_amount"]),
            unmatched_amount=max(Decimal(bucket["net_amount"]) - Decimal(bucket["matched_amount"]), Decimal("0.00")),
            reconciliation_rate=decimal_rate(Decimal(bucket["matched_amount"]), Decimal(bucket["net_amount"])),
            record_count=int(bucket["record_count"]),
        )
        for channel, bucket in buckets.items()
    ]
    breakdown.sort(key=lambda item: item.gross_amount, reverse=True)
    return breakdown


def apply_period_range(query, column, period_start: str | None, period_end: str | None):
    if period_start:
        query = query.where(column >= period_start)
    if period_end:
        query = query.where(column <= period_end)
    return query


def parse_period_start(period: str) -> datetime:
    try:
        return datetime.strptime(period, "%Y-%m")
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="Invalid period format, expected YYYY-MM") from exc


def next_period_start(period: str) -> datetime:
    current = parse_period_start(period)
    if current.month == 12:
        return current.replace(year=current.year + 1, month=1)
    return current.replace(month=current.month + 1)


def apply_datetime_period_range(query, column, period_start: str | None, period_end: str | None):
    if period_start:
        query = query.where(column >= parse_period_start(period_start))
    if period_end:
        query = query.where(column < next_period_start(period_end))
    return query


def source_document_filter(document_ids: list[str]):
    return or_(
        ExpenseItem.source_document_id.in_(document_ids),
        *(ExpenseItem.source_document_id.startswith(f"{document_id}:", autoescape=True) for document_id in document_ids),
    )


def empty_financial_analytics_report() -> FinancialAnalyticsReport:
    return FinancialAnalyticsReport(
        metrics=FinancialAnalyticsMetrics(
            store_count=0,
            period_count=0,
            total_income_amount=Decimal("0.00"),
            total_net_income_amount=Decimal("0.00"),
            total_fee_amount=Decimal("0.00"),
            approval_month_amount=Decimal("0.00"),
            total_expense_amount=Decimal("0.00"),
            total_profit_amount=Decimal("0.00"),
            bank_expense_amount=Decimal("0.00"),
            matched_expense_amount=Decimal("0.00"),
            unmatched_bank_amount=Decimal("0.00"),
            unmatched_bank_count=0,
            pending_expense_amount=Decimal("0.00"),
            pending_expense_count=0,
            confirmed_match_count=0,
            pending_match_count=0,
            bank_not_occurred_count=0,
        ),
        trends=[],
        stores=[],
        revenue_channels=[],
        revenue_channel_monthly_summary=[],
        approval_status_summary=[],
        categories=[],
        templates=[],
        reconciliation=[],
    )


def authorized_report_store_ids(
    session: Session,
    shareholder_grant: ShareholderAccessGrant | None,
    current_user: User | None,
    store_id: str | None = None,
) -> list[str]:
    stores_query = select(Store).order_by(Store.created_at.desc())
    if store_id:
        ensure_report_store_access(session, shareholder_grant, current_user, store_id)
        stores_query = stores_query.where(Store.id == store_id)
    stores = filter_report_stores(session, list(session.scalars(stores_query)), shareholder_grant, current_user)
    return [store.id for store in stores]


@router.get("/analytics", response_model=ApiEnvelope[FinancialAnalyticsReport])
def read_financial_analytics(
    store_id: str | None = None,
    period_start: str | None = None,
    period_end: str | None = None,
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> ApiEnvelope[FinancialAnalyticsReport]:
    store_ids = authorized_report_store_ids(session, shareholder_grant, current_user, store_id)
    stores = list(session.scalars(select(Store).where(Store.id.in_(store_ids)))) if store_ids else []
    if not store_ids:
        return ApiEnvelope(data=empty_financial_analytics_report())

    revenue_query = select(RevenueRecord).where(RevenueRecord.store_id.in_(store_ids))
    revenue_query = apply_period_range(revenue_query, RevenueRecord.ledger_period, period_start, period_end)
    expense_query = select(ExpenseItem).where(ExpenseItem.store_id.in_(store_ids))
    expense_query = apply_period_range(expense_query, ExpenseItem.ledger_period, period_start, period_end)
    bank_query = select(BankTransaction).where(BankTransaction.store_id.in_(store_ids))
    bank_query = apply_period_range(bank_query, BankTransaction.ledger_period, period_start, period_end)

    revenue_rows = session.scalars(revenue_query).all()
    expense_rows = session.scalars(expense_query).all()
    bank_rows = session.scalars(bank_query).all()
    revenue_match_rows = session.execute(
        select(RevenueBankMatch, RevenueRecord)
        .join(RevenueBankMatchRecord, RevenueBankMatchRecord.revenue_bank_match_id == RevenueBankMatch.id)
        .join(RevenueRecord, RevenueBankMatchRecord.revenue_record_id == RevenueRecord.id)
        .where(RevenueRecord.store_id.in_(store_ids))
    ).all()
    revenue_match_rows = [
        (match, record)
        for match, record in revenue_match_rows
        if (not period_start or (match.accounting_period or record.ledger_period or "") >= period_start)
        and (not period_end or (match.accounting_period or record.ledger_period or "") <= period_end)
    ]
    match_rows = session.execute(
        select(ExpenseBankMatch, ExpenseItem)
        .join(ExpenseItem, ExpenseBankMatch.expense_item_id == ExpenseItem.id)
        .where(ExpenseItem.store_id.in_(store_ids))
    ).all()
    match_rows = [
        (match, expense)
        for match, expense in match_rows
        if (not period_start or (match.accounting_period or expense.ledger_period or "") >= period_start)
        and (not period_end or (match.accounting_period or expense.ledger_period or "") <= period_end)
    ]

    total_income_amount = sum((Decimal(record.gross_amount) for record in revenue_rows), Decimal("0.00"))
    total_net_income_amount = sum((Decimal(record.net_amount) for record in revenue_rows), Decimal("0.00"))
    total_fee_amount = sum((Decimal(record.fee_amount) for record in revenue_rows), Decimal("0.00"))
    total_expense_amount = sum((Decimal(item.amount) for item in expense_rows), Decimal("0.00"))
    bank_expense_amount = sum(
        (Decimal(transaction.amount) for transaction in bank_rows if transaction.direction == "expense"),
        Decimal("0.00"),
    )
    matched_expense_amount = sum(
        (Decimal(match.amount) for match, _expense in match_rows if match.status == MatchStatus.CONFIRMED.value),
        Decimal("0.00"),
    )
    unmatched_bank_rows = [
        transaction
        for transaction in bank_rows
        if transaction.direction == "expense" and Decimal(transaction.amount) > Decimal(transaction.matched_amount or 0)
    ]
    unmatched_bank_amount = sum(
        (Decimal(transaction.amount) - Decimal(transaction.matched_amount or 0) for transaction in unmatched_bank_rows),
        Decimal("0.00"),
    )
    pending_expense_rows = [
        item
        for item in expense_rows
        if item.payment_status in {"unpaid", "partial_paid", "no_bank_flow"}
    ]
    pending_expense_amount = sum((Decimal(item.amount) for item in pending_expense_rows), Decimal("0.00"))

    periods = sorted(
        {
            *(record.ledger_period for record in revenue_rows if record.ledger_period),
            *(item.ledger_period for item in expense_rows if item.ledger_period),
            *(transaction.ledger_period for transaction in bank_rows if transaction.ledger_period),
        }
    )
    trend_items: list[FinancialAnalyticsTrendItem] = []
    for period in periods[-12:]:
        period_income = sum(
            (Decimal(record.gross_amount) for record in revenue_rows if record.ledger_period == period),
            Decimal("0.00"),
        )
        period_expense = sum(
            (Decimal(item.amount) for item in expense_rows if item.ledger_period == period),
            Decimal("0.00"),
        )
        period_bank_expense = sum(
            (
                Decimal(transaction.amount)
                for transaction in bank_rows
                if transaction.ledger_period == period and transaction.direction == "expense"
            ),
            Decimal("0.00"),
        )
        period_matched = sum(
            (
                Decimal(match.amount)
                for match, expense in match_rows
                if match.status == MatchStatus.CONFIRMED.value
                and (match.accounting_period or expense.ledger_period) == period
            ),
            Decimal("0.00"),
        )
        period_unmatched = sum(
            (
                Decimal(transaction.amount) - Decimal(transaction.matched_amount or 0)
                for transaction in unmatched_bank_rows
                if transaction.ledger_period == period
            ),
            Decimal("0.00"),
        )
        period_bank_count = sum(1 for transaction in bank_rows if transaction.ledger_period == period)
        trend_items.append(
            FinancialAnalyticsTrendItem(
                period=period,
                income_amount=period_income,
                expense_amount=period_expense,
                profit_amount=period_income - period_expense,
                bank_expense_amount=period_bank_expense,
                matched_expense_amount=period_matched,
                unmatched_bank_amount=period_unmatched,
                bank_transaction_count=period_bank_count,
            )
        )

    store_items: list[FinancialAnalyticsStoreItem] = []
    for store in stores:
        store_income = sum(
            (Decimal(record.gross_amount) for record in revenue_rows if record.store_id == store.id),
            Decimal("0.00"),
        )
        store_expense = sum(
            (Decimal(item.amount) for item in expense_rows if item.store_id == store.id),
            Decimal("0.00"),
        )
        store_bank_expense = sum(
            (
                Decimal(transaction.amount)
                for transaction in bank_rows
                if transaction.store_id == store.id and transaction.direction == "expense"
            ),
            Decimal("0.00"),
        )
        store_matched = sum(
            (
                Decimal(match.amount)
                for match, expense in match_rows
                if expense.store_id == store.id and match.status == MatchStatus.CONFIRMED.value
            ),
            Decimal("0.00"),
        )
        store_unmatched_rows = [transaction for transaction in unmatched_bank_rows if transaction.store_id == store.id]
        store_pending_expense_rows = [item for item in pending_expense_rows if item.store_id == store.id]
        store_items.append(
            FinancialAnalyticsStoreItem(
                store_id=store.id,
                store_name=store.name,
                income_amount=store_income,
                expense_amount=store_expense,
                profit_amount=store_income - store_expense,
                bank_expense_amount=store_bank_expense,
                matched_expense_amount=store_matched,
                unmatched_bank_amount=sum(
                    (Decimal(item.amount) - Decimal(item.matched_amount or 0) for item in store_unmatched_rows),
                    Decimal("0.00"),
                ),
                unmatched_bank_count=len(store_unmatched_rows),
                pending_expense_amount=sum((Decimal(item.amount) for item in store_pending_expense_rows), Decimal("0.00")),
                pending_expense_count=len(store_pending_expense_rows),
            )
        )
    store_items.sort(key=lambda item: item.expense_amount, reverse=True)

    category_bucket: dict[tuple[str, str | None], tuple[Decimal, int]] = {}
    for item in expense_rows:
        key = (item.category_l1 or "未分类", item.category_l2)
        amount, count = category_bucket.get(key, (Decimal("0.00"), 0))
        category_bucket[key] = (amount + Decimal(item.amount), count + 1)
    category_items = [
        FinancialAnalyticsCategoryItem(category_l1=key[0], category_l2=key[1], amount=value[0], item_count=value[1])
        for key, value in category_bucket.items()
    ]
    category_items.sort(key=lambda item: item.amount, reverse=True)
    revenue_channel_monthly_bucket: dict[tuple[str, str], dict[str, Decimal | int]] = {}
    for record in revenue_rows:
        if not record.ledger_period:
            continue
        bucket = revenue_channel_monthly_bucket.setdefault(
            (record.ledger_period, record.channel),
            {
                "gross_amount": Decimal("0.00"),
                "net_amount": Decimal("0.00"),
                "record_count": 0,
            },
        )
        bucket["gross_amount"] = Decimal(bucket["gross_amount"]) + Decimal(record.gross_amount)
        bucket["net_amount"] = Decimal(bucket["net_amount"]) + Decimal(record.net_amount)
        bucket["record_count"] = int(bucket["record_count"]) + 1
    revenue_channel_monthly_summary = [
        RevenueChannelMonthlyBreakdownItem(
            period=period,
            channel=channel,
            gross_amount=Decimal(bucket["gross_amount"]),
            net_amount=Decimal(bucket["net_amount"]),
            record_count=int(bucket["record_count"]),
        )
        for (period, channel), bucket in sorted(
            revenue_channel_monthly_bucket.items(),
            key=lambda item: (item[0][0], Decimal(item[1]["gross_amount"])),
        )
    ]
    revenue_channel_breakdown = build_revenue_channel_breakdown(
        revenue_rows,
        [match for match, _record in revenue_match_rows],
    )

    approval_instances_query = (
        select(ApprovalInstance, ApprovalTemplate)
        .join(ApprovalTemplate, ApprovalInstance.template_id == ApprovalTemplate.id)
        .where(ApprovalInstance.store_id.in_(store_ids))
    )
    approval_instances_query = apply_datetime_period_range(
        approval_instances_query,
        ApprovalInstance.submit_at,
        period_start,
        period_end,
    )
    approval_instances = session.execute(approval_instances_query).all()
    approval_stats_map = approval_expense_stats_map(session, [approval.id for approval, _template in approval_instances])
    approvals_by_document_id = {
        approval.dingtalk_instance_id: (approval, template)
        for approval, template in approval_instances
    }
    approval_month_amount = sum(
        (
            Decimal(approval_stats_map.get(approval.id, {}).get("total_expense_amount", Decimal("0.00")))
            for approval, _template in approval_instances
        ),
        Decimal("0.00"),
    )
    approval_status_bucket: dict[str, dict[str, Decimal | int]] = {}
    for approval, _template in approval_instances:
        stats = approval_stats_map.get(approval.id, {})
        bucket = approval_status_bucket.setdefault(
            approval.processing_status or "unparsed",
            {"count": 0, "amount": Decimal("0.00")},
        )
        bucket["count"] = int(bucket["count"]) + 1
        bucket["amount"] = Decimal(bucket["amount"]) + Decimal(stats.get("total_expense_amount", Decimal("0.00")))
    template_approval_counts: dict[str, tuple[str, int]] = {}
    for approval, template in approval_instances:
        count = template_approval_counts.get(template.id, (template.name, 0))[1]
        template_approval_counts[template.id] = (template.name, count + 1)

    expense_ids_by_template: dict[str, set[str]] = {}
    expense_amount_by_template: dict[str, Decimal] = {}
    for item in expense_rows:
        document_id = item.source_document_id.split(":", 1)[0] if item.source_document_id else None
        if not document_id or document_id not in approvals_by_document_id:
            continue
        _approval, template = approvals_by_document_id[document_id]
        expense_ids_by_template.setdefault(template.id, set()).add(item.id)
        expense_amount_by_template[template.id] = expense_amount_by_template.get(template.id, Decimal("0.00")) + Decimal(item.amount)

    matched_amount_by_template: dict[str, Decimal] = {}
    for match, expense in match_rows:
        if match.status != MatchStatus.CONFIRMED.value:
            continue
        for current_template_id, expense_ids in expense_ids_by_template.items():
            if expense.id in expense_ids:
                matched_amount_by_template[current_template_id] = (
                    matched_amount_by_template.get(current_template_id, Decimal("0.00")) + Decimal(match.amount)
                )
                break

    template_items: list[FinancialAnalyticsTemplateItem] = []
    for template_id, (template_name, approval_count) in template_approval_counts.items():
        template_items.append(
            FinancialAnalyticsTemplateItem(
                template_id=template_id,
                template_name=template_name,
                approval_count=approval_count,
                expense_amount=expense_amount_by_template.get(template_id, Decimal("0.00")),
                matched_amount=matched_amount_by_template.get(template_id, Decimal("0.00")),
            )
        )
    template_items.sort(key=lambda item: item.approval_count, reverse=True)

    reconciliation_items = [
        FinancialAnalyticsReconciliationItem(
            status=status,
            count=sum(1 for match, _expense in match_rows if match.status == status),
            amount=sum(
                (Decimal(match.amount) for match, _expense in match_rows if match.status == status),
                Decimal("0.00"),
            ),
        )
        for status in (MatchStatus.CONFIRMED.value, MatchStatus.CANDIDATE.value, MatchStatus.REJECTED.value)
    ]
    approval_status_summary = [
        FinancialAnalyticsApprovalStatusItem(
            status=status,
            count=int(bucket["count"]),
            amount=Decimal(bucket["amount"]),
        )
        for status, bucket in approval_status_bucket.items()
    ]
    approval_status_summary.sort(key=lambda item: item.count, reverse=True)

    report = FinancialAnalyticsReport(
        metrics=FinancialAnalyticsMetrics(
            store_count=len(stores),
            period_count=len(periods),
            total_income_amount=total_income_amount,
            total_net_income_amount=total_net_income_amount,
            total_fee_amount=total_fee_amount,
            approval_month_amount=approval_month_amount,
            total_expense_amount=total_expense_amount,
            total_profit_amount=total_income_amount - total_expense_amount,
            bank_expense_amount=bank_expense_amount,
            matched_expense_amount=matched_expense_amount,
            unmatched_bank_amount=unmatched_bank_amount,
            unmatched_bank_count=len(unmatched_bank_rows),
            pending_expense_amount=pending_expense_amount,
            pending_expense_count=len(pending_expense_rows),
            confirmed_match_count=sum(1 for match, _expense in match_rows if match.status == MatchStatus.CONFIRMED.value),
            pending_match_count=sum(1 for match, _expense in match_rows if match.status == MatchStatus.CANDIDATE.value),
            bank_not_occurred_count=sum(
                1 for match, _expense in match_rows if match.status == MatchStatus.CONFIRMED.value and not match.bank_occurred
            ),
        ),
        trends=trend_items,
        stores=store_items[:20],
        revenue_channels=revenue_channel_breakdown[:20],
        revenue_channel_monthly_summary=revenue_channel_monthly_summary,
        approval_status_summary=approval_status_summary,
        categories=category_items[:20],
        templates=template_items[:20],
        reconciliation=reconciliation_items,
    )
    return ApiEnvelope(data=report)


@router.get("/analytics/details", response_model=ApiEnvelope[FinancialAnalyticsDetailReport])
def read_financial_analytics_details(
    detail_type: str,
    store_id: str | None = None,
    period_start: str | None = None,
    period_end: str | None = None,
    category_l1: str | None = None,
    category_l2: str | None = None,
    template_id: str | None = None,
    match_status: str | None = None,
    page_size: int = 100,
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> ApiEnvelope[FinancialAnalyticsDetailReport]:
    page_size = min(max(page_size, 1), 500)
    store_ids = authorized_report_store_ids(session, shareholder_grant, current_user, store_id)
    if not store_ids:
        return ApiEnvelope(data=FinancialAnalyticsDetailReport(title="暂无授权门店"))

    expense_query = select(ExpenseItem).where(ExpenseItem.store_id.in_(store_ids))
    expense_query = apply_period_range(expense_query, ExpenseItem.ledger_period, period_start, period_end)
    bank_query = select(BankTransaction).where(BankTransaction.store_id.in_(store_ids))
    bank_query = apply_period_range(bank_query, BankTransaction.ledger_period, period_start, period_end)
    approval_query = select(ApprovalInstance).where(ApprovalInstance.store_id.in_(store_ids))
    approval_query = apply_datetime_period_range(approval_query, ApprovalInstance.submit_at, period_start, period_end)

    title = "统计明细"
    expense_items: list[ExpenseItem] = []
    bank_transactions: list[BankTransaction] = []
    approval_instances: list[ApprovalInstance] = []
    reconciliation_records: list[ReconciliationRecord] = []

    if detail_type == "unmatched_bank":
        title = "未对账银行流水"
        bank_transactions = list(
            session.scalars(
                bank_query.where(
                    BankTransaction.direction == "expense",
                    BankTransaction.matched_amount < BankTransaction.amount,
                )
                .order_by(BankTransaction.occurred_at.desc())
                .limit(page_size)
            )
        )
    elif detail_type == "pending_expense":
        title = "未付款审批支出"
        expense_items = list(
            session.scalars(
                expense_query.where(ExpenseItem.payment_status.in_(["unpaid", "partial_paid", "no_bank_flow"]))
                .order_by(ExpenseItem.expense_date.desc().nullslast(), ExpenseItem.created_at.desc())
                .limit(page_size)
            )
        )
    elif detail_type == "store":
        title = "门店经营明细"
        expense_items = list(
            session.scalars(
                expense_query.order_by(ExpenseItem.expense_date.desc().nullslast(), ExpenseItem.created_at.desc()).limit(page_size)
            )
        )
        bank_transactions = list(
            session.scalars(bank_query.order_by(BankTransaction.occurred_at.desc()).limit(page_size))
        )
        approval_instances = list(
            session.scalars(approval_query.order_by(ApprovalInstance.created_at.desc()).limit(page_size))
        )
    elif detail_type == "category":
        title = "费用分类明细"
        if category_l1:
            expense_query = expense_query.where(ExpenseItem.category_l1 == category_l1)
        if category_l2:
            expense_query = expense_query.where(ExpenseItem.category_l2 == category_l2)
        expense_items = list(
            session.scalars(
                expense_query.order_by(ExpenseItem.expense_date.desc().nullslast(), ExpenseItem.created_at.desc()).limit(page_size)
            )
        )
    elif detail_type == "template":
        title = "审批模版明细"
        if not template_id:
            raise HTTPException(status_code=422, detail="template_id is required")
        approval_instances = list(
            session.scalars(
                approval_query.where(ApprovalInstance.template_id == template_id)
                .order_by(ApprovalInstance.created_at.desc())
                .limit(page_size)
            )
        )
        document_ids = [instance.dingtalk_instance_id for instance in approval_instances]
        if document_ids:
            expense_items = list(
                session.scalars(
                    expense_query.where(source_document_filter(document_ids))
                    .order_by(ExpenseItem.expense_date.desc().nullslast(), ExpenseItem.created_at.desc())
                    .limit(page_size)
                )
            )
    elif detail_type == "reconciliation":
        title = "对账记录明细"
        target_status = match_status or MatchStatus.CONFIRMED.value
        rows = session.execute(
            select(ExpenseBankMatch, BankTransaction, ExpenseItem, ApprovalInstance, ApprovalTemplate)
            .join(BankTransaction, ExpenseBankMatch.bank_transaction_id == BankTransaction.id)
            .join(ExpenseItem, ExpenseBankMatch.expense_item_id == ExpenseItem.id)
            .outerjoin(
                ApprovalInstance,
                (ExpenseItem.source_document_id == ApprovalInstance.dingtalk_instance_id)
                | (
                    func.substr(
                        ExpenseItem.source_document_id,
                        1,
                        func.length(ApprovalInstance.dingtalk_instance_id) + 1,
                    )
                    == ApprovalInstance.dingtalk_instance_id + ":"
                ),
            )
            .outerjoin(ApprovalTemplate, ApprovalInstance.template_id == ApprovalTemplate.id)
            .where(ExpenseItem.store_id.in_(store_ids), ExpenseBankMatch.status == target_status)
            .order_by(ExpenseBankMatch.created_at.desc())
            .limit(page_size)
        ).all()
        reconciliation_records = [
            ReconciliationRecord(
                match=match,
                bank_transaction=bank_transaction,
                expense_item=expense,
                approval_instance=approval,
                template_name=template.name if template else None,
                display_fields={},
            )
            for match, bank_transaction, expense, approval, template in rows
        ]
    else:
        raise HTTPException(status_code=422, detail="Unknown detail type")

    return ApiEnvelope(
        data=FinancialAnalyticsDetailReport(
            title=title,
            expense_items=expense_items,
            bank_transactions=bank_transactions,
            approval_instances=approval_instances,
            reconciliation_records=reconciliation_records,
        )
    )


@router.get("/store-summaries", response_model=ApiEnvelope[list[StoreReportSummary]])
def list_store_summaries(
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> ApiEnvelope[list[StoreReportSummary]]:
    stores = session.scalars(select(Store).order_by(Store.created_at.desc())).all()
    stores = filter_report_stores(session, list(stores), shareholder_grant, current_user)
    summaries: list[StoreReportSummary] = []
    for store in stores:
        ledger_query = select(Ledger).where(Ledger.store_id == store.id)
        if shareholder_grant is not None:
            ledger_query = ledger_query.where(Ledger.status == "closed")
        ledger = session.scalar(
            ledger_query.order_by(Ledger.period.desc(), Ledger.created_at.desc())
        )
        if ledger is None:
            continue
        summary = build_report_summary(session, ledger, store)
        summaries.append(StoreReportSummary(**summary.model_dump(), ledger_id=ledger.id))
    return ApiEnvelope(data=summaries)


@router.get("/periods", response_model=ApiEnvelope[list[ReportPeriodOption]])
def list_report_periods(
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> ApiEnvelope[list[ReportPeriodOption]]:
    allowed_store_ids: set[str] | None = None
    if shareholder_grant is not None:
        allowed_store_ids = set(grant_store_ids(shareholder_grant))
    elif current_user is not None:
        allowed_store_ids = set(effective_store_ids(session, current_user))

    query = select(Ledger.period, func.count(Ledger.store_id)).group_by(Ledger.period)
    if shareholder_grant is not None:
        query = query.where(Ledger.status == "closed")
    if allowed_store_ids is not None:
        query = query.where(Ledger.store_id.in_(allowed_store_ids))
    rows = session.execute(query.order_by(Ledger.period.desc())).all()
    return ApiEnvelope(
        data=[ReportPeriodOption(period=period, store_count=store_count) for period, store_count in rows]
    )


@router.get("/ledger-periods", response_model=ApiEnvelope[list[LedgerPeriodOption]])
def list_ledger_periods(
    store_id: str,
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> ApiEnvelope[list[LedgerPeriodOption]]:
    ensure_report_store_access(session, shareholder_grant, current_user, store_id)
    store = session.get(Store, store_id)
    if store is None:
        raise HTTPException(status_code=404, detail="Store not found")
    query = select(Ledger).where(Ledger.store_id == store_id)
    if shareholder_grant is not None:
        query = query.where(Ledger.status == "closed")
    ledgers = session.scalars(query.order_by(Ledger.period.desc(), Ledger.created_at.desc())).all()
    return ApiEnvelope(
        data=[
            LedgerPeriodOption(
                ledger_id=ledger.id,
                store_id=ledger.store_id,
                period=ledger.period,
                ledger_status=ledger.status,
            )
            for ledger in ledgers
        ]
    )


@router.get("/ledger-trends", response_model=ApiEnvelope[list[LedgerTrend]])
def list_ledger_trends(
    store_id: str | None = None,
    limit: int = 6,
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> ApiEnvelope[list[LedgerTrend]]:
    limit = min(max(limit, 1), 24)
    stores_query = select(Store).order_by(Store.created_at.desc())
    if store_id:
        ensure_report_store_access(session, shareholder_grant, current_user, store_id)
        stores_query = stores_query.where(Store.id == store_id)
    stores = session.scalars(stores_query).all()
    stores = filter_report_stores(session, list(stores), shareholder_grant, current_user)

    trends: list[LedgerTrend] = []
    for store in stores:
        ledger_query = select(Ledger).where(Ledger.store_id == store.id)
        if shareholder_grant is not None:
            ledger_query = ledger_query.where(Ledger.status == "closed")
        ledgers = session.scalars(
            ledger_query.order_by(Ledger.period.desc(), Ledger.created_at.desc()).limit(limit)
        ).all()
        if not ledgers:
            continue
        summaries = [build_report_summary(session, ledger, store) for ledger in reversed(ledgers)]
        trends.append(LedgerTrend(store_id=store.id, store_name=store.name, items=summaries))
    return ApiEnvelope(data=trends)


@router.get("/store-comparison", response_model=ApiEnvelope[StoreComparisonReport])
def read_store_comparison(
    period: str | None = None,
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> ApiEnvelope[StoreComparisonReport]:
    allowed_store_ids: set[str] | None = None
    if shareholder_grant is not None:
        allowed_store_ids = set(grant_store_ids(shareholder_grant))
    elif current_user is not None:
        allowed_store_ids = set(effective_store_ids(session, current_user))

    period_query = select(func.max(Ledger.period))
    if shareholder_grant is not None:
        period_query = period_query.where(Ledger.status == "closed")
    if allowed_store_ids is not None:
        period_query = period_query.where(Ledger.store_id.in_(allowed_store_ids))
    target_period = period or session.scalar(period_query)
    if not target_period:
        raise HTTPException(status_code=404, detail="No ledger period available")

    ledger_query = select(Ledger).where(Ledger.period == target_period)
    if shareholder_grant is not None:
        ledger_query = ledger_query.where(Ledger.status == "closed")
    if allowed_store_ids is not None:
        ledger_query = ledger_query.where(Ledger.store_id.in_(allowed_store_ids))
    ledgers = session.scalars(ledger_query.order_by(Ledger.created_at.desc())).all()
    summaries: list[LedgerReportSummary] = []
    for ledger in ledgers:
        store = session.get(Store, ledger.store_id)
        if store is None:
            continue
        summaries.append(build_report_summary(session, ledger, store))
    summaries.sort(key=lambda item: item.profit_amount, reverse=True)
    return ApiEnvelope(
        data=StoreComparisonReport(
            period=target_period,
            items=summaries,
            total_income_amount=sum((item.income_amount for item in summaries), Decimal("0.00")),
            total_expense_amount=sum((item.expense_amount for item in summaries), Decimal("0.00")),
            total_profit_amount=sum((item.profit_amount for item in summaries), Decimal("0.00")),
        )
    )


@router.get("/ledger-summary", response_model=ApiEnvelope[LedgerReportSummary])
def read_ledger_summary(
    store_id: str,
    period: str,
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> ApiEnvelope[LedgerReportSummary]:
    store, ledger = read_ledger_for_report(session, store_id, period, shareholder_grant, current_user)
    return ApiEnvelope(data=build_report_summary(session, ledger, store))


@router.get("/ledger-detail", response_model=ApiEnvelope[LedgerReportDetail])
def read_ledger_detail(
    store_id: str,
    period: str,
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> ApiEnvelope[LedgerReportDetail]:
    store, ledger = read_ledger_for_report(session, store_id, period, shareholder_grant, current_user)

    category_name = func.coalesce(ExpenseItem.category_l1, "未分类")
    supplier_name = func.coalesce(ExpenseItem.supplier_name, "未关联供应商")
    category_rows = session.execute(
        select(
            category_name,
            func.sum(ExpenseItem.amount),
            func.count(),
        )
        .where(ExpenseItem.store_id == store_id, ExpenseItem.ledger_period == period)
        .group_by(category_name)
        .order_by(func.sum(ExpenseItem.amount).desc())
    ).all()
    supplier_rows = session.execute(
        select(
            supplier_name,
            func.sum(ExpenseItem.amount),
            func.count(),
        )
        .where(ExpenseItem.store_id == store_id, ExpenseItem.ledger_period == period)
        .group_by(supplier_name)
        .order_by(func.sum(ExpenseItem.amount).desc())
    ).all()
    pending_expense_items = session.scalars(
        select(ExpenseItem)
        .where(
            ExpenseItem.store_id == store_id,
            ExpenseItem.ledger_period == period,
            ExpenseItem.payment_status.in_(["unpaid", "partial_paid"]),
        )
        .order_by(ExpenseItem.created_at.desc())
        .limit(20)
    ).all()
    pending_bank_transactions = session.scalars(
        select(BankTransaction)
        .where(
            BankTransaction.store_id == store_id,
            BankTransaction.ledger_period == period,
            BankTransaction.matched_amount < BankTransaction.amount,
        )
        .order_by(BankTransaction.occurred_at.desc())
        .limit(20)
    ).all()
    revenue_records = session.scalars(
        select(RevenueRecord)
        .where(RevenueRecord.store_id == store_id, RevenueRecord.ledger_period == period)
        .order_by(RevenueRecord.revenue_date.desc(), RevenueRecord.created_at.desc())
        .limit(50)
    ).all()
    all_revenue_records = list(
        session.scalars(
            select(RevenueRecord)
            .where(RevenueRecord.store_id == store_id, RevenueRecord.ledger_period == period)
            .order_by(RevenueRecord.revenue_date.asc(), RevenueRecord.created_at.asc())
        )
    )
    revenue_matches = list(
        session.scalars(
            select(RevenueBankMatch)
            .join(BankTransaction, RevenueBankMatch.bank_transaction_id == BankTransaction.id)
            .where(BankTransaction.store_id == store_id, BankTransaction.ledger_period == period)
        )
    )

    return ApiEnvelope(
        data=LedgerReportDetail(
            summary=build_report_summary(session, ledger, store),
            revenue_records=list(revenue_records),
            revenue_channel_breakdown=build_revenue_channel_breakdown(all_revenue_records, revenue_matches),
            category_breakdown=[
                ExpenseBreakdownItem(name=name, amount=decimal_sum(amount), item_count=count)
                for name, amount, count in category_rows
            ],
            supplier_breakdown=[
                ExpenseBreakdownItem(name=name, amount=decimal_sum(amount), item_count=count)
                for name, amount, count in supplier_rows
            ],
            pending_expense_items=list(pending_expense_items),
            pending_bank_transactions=list(pending_bank_transactions),
        )
    )


@router.get("/ledger-detail.csv")
def export_ledger_detail_csv(
    store_id: str,
    period: str,
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> StreamingResponse:
    store, ledger = read_ledger_for_report(session, store_id, period, shareholder_grant, current_user)
    revenue_records, expenses, bank_transactions = read_ledger_export_rows(session, store_id, period)
    revenue_matches = list(
        session.scalars(
            select(RevenueBankMatch)
            .join(BankTransaction, RevenueBankMatch.bank_transaction_id == BankTransaction.id)
            .where(BankTransaction.store_id == store_id, BankTransaction.ledger_period == period)
        )
    )
    channel_breakdown = build_revenue_channel_breakdown(revenue_records, revenue_matches)
    summary = build_report_summary(session, ledger, store)

    buffer = io.StringIO()
    buffer.write("\ufeff")
    writer = csv.writer(buffer)
    writer.writerow(["报表类型", "门店", "账期", "账套状态", "收入", "支出", "利润"])
    writer.writerow(
        [
            "账套汇总",
            summary.store_name,
            summary.period,
            summary.ledger_status,
            summary.income_amount,
            summary.expense_amount,
            summary.profit_amount,
        ]
    )
    writer.writerow([])
    writer.writerow(["营业收入"])
    writer.writerow(["日期", "渠道", "经营收入", "实收金额", "手续费", "备注"])
    for record in revenue_records:
        writer.writerow(
            [
                record.revenue_date.isoformat(),
                record.channel,
                record.gross_amount,
                record.net_amount,
                record.fee_amount,
                record.remark or "",
            ]
        )
    writer.writerow([])
    writer.writerow(["收入渠道汇总"])
    writer.writerow(["渠道", "经营收入", "实收金额", "手续费", "费率", "已对账", "未对账", "对账完成率", "记录数"])
    for item in channel_breakdown:
        writer.writerow(
            [
                item.channel,
                item.gross_amount,
                item.net_amount,
                item.fee_amount,
                f"{item.fee_rate}%",
                item.matched_amount,
                item.unmatched_amount,
                f"{item.reconciliation_rate}%",
                item.record_count,
            ]
        )
    writer.writerow([])
    writer.writerow(["支出明细"])
    writer.writerow(["日期", "说明", "金额", "一级分类", "二级分类", "供应商", "收款账号", "付款状态"])
    for item in expenses:
        writer.writerow(
            [
                item.expense_date.isoformat() if item.expense_date else "",
                item.description,
                item.amount,
                item.category_l1 or "",
                item.category_l2 or "",
                item.supplier_name or "",
                item.payee_account or "",
                item.payment_status,
            ]
        )
    writer.writerow([])
    writer.writerow(["银行流水"])
    writer.writerow(["发生时间", "方向", "金额", "对方户名", "对方账号", "摘要", "流水号", "已匹配金额"])
    for transaction in bank_transactions:
        writer.writerow(
            [
                transaction.occurred_at.isoformat(sep=" "),
                transaction.direction,
                transaction.amount,
                transaction.counterparty_name or "",
                transaction.counterparty_account or "",
                transaction.summary or "",
                transaction.bank_serial_no or "",
                transaction.matched_amount,
            ]
        )
    buffer.seek(0)
    filename = f"ledger-detail-{store_id}-{period}.csv"
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def append_rows(sheet, rows: list[list[object]]) -> None:
    for row in rows:
        sheet.append(row)


@router.get("/ledger-detail.xlsx")
def export_ledger_detail_xlsx(
    store_id: str,
    period: str,
    session: Session = Depends(get_session),
    shareholder_grant: ShareholderAccessGrant | None = Depends(get_report_access),
    current_user: User | None = Depends(get_optional_current_user),
) -> StreamingResponse:
    store, ledger = read_ledger_for_report(session, store_id, period, shareholder_grant, current_user)
    revenue_records, expenses, bank_transactions = read_ledger_export_rows(session, store_id, period)
    revenue_matches = list(
        session.scalars(
            select(RevenueBankMatch)
            .join(BankTransaction, RevenueBankMatch.bank_transaction_id == BankTransaction.id)
            .where(BankTransaction.store_id == store_id, BankTransaction.ledger_period == period)
        )
    )
    channel_breakdown = build_revenue_channel_breakdown(revenue_records, revenue_matches)
    summary = build_report_summary(session, ledger, store)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "账套汇总"
    append_rows(
        summary_sheet,
        [
            ["门店", summary.store_name],
            ["账期", summary.period],
            ["账套状态", summary.ledger_status],
            ["收入", decimal_cell(summary.income_amount)],
            ["支出", decimal_cell(summary.expense_amount)],
            ["利润", decimal_cell(summary.profit_amount)],
            ["待付款支出", summary.pending_expense_count],
            ["未匹配流水", summary.pending_bank_transaction_count],
        ],
    )

    revenue_sheet = workbook.create_sheet("营业收入")
    revenue_sheet.append(["日期", "渠道", "经营收入", "实收金额", "手续费", "备注"])
    for record in revenue_records:
        revenue_sheet.append(
            [
                record.revenue_date.isoformat(),
                record.channel,
                decimal_cell(record.gross_amount),
                decimal_cell(record.net_amount),
                decimal_cell(record.fee_amount),
                record.remark or "",
            ]
        )

    revenue_channel_sheet = workbook.create_sheet("收入渠道汇总")
    revenue_channel_sheet.append(["渠道", "经营收入", "实收金额", "手续费", "费率", "已对账", "未对账", "对账完成率", "记录数"])
    for item in channel_breakdown:
        revenue_channel_sheet.append(
            [
                item.channel,
                decimal_cell(item.gross_amount),
                decimal_cell(item.net_amount),
                decimal_cell(item.fee_amount),
                float(item.fee_rate),
                decimal_cell(item.matched_amount),
                decimal_cell(item.unmatched_amount),
                float(item.reconciliation_rate),
                item.record_count,
            ]
        )

    expense_sheet = workbook.create_sheet("支出明细")
    expense_sheet.append(["日期", "说明", "金额", "一级分类", "二级分类", "供应商", "收款账号", "付款状态"])
    for item in expenses:
        expense_sheet.append(
            [
                item.expense_date.isoformat() if item.expense_date else "",
                item.description,
                decimal_cell(item.amount),
                item.category_l1 or "",
                item.category_l2 or "",
                item.supplier_name or "",
                item.payee_account or "",
                item.payment_status,
            ]
        )

    bank_sheet = workbook.create_sheet("银行流水")
    bank_sheet.append(["发生时间", "方向", "金额", "对方户名", "对方账号", "摘要", "流水号", "已匹配金额"])
    for transaction in bank_transactions:
        bank_sheet.append(
            [
                transaction.occurred_at.isoformat(sep=" "),
                transaction.direction,
                decimal_cell(transaction.amount),
                transaction.counterparty_name or "",
                transaction.counterparty_account or "",
                transaction.summary or "",
                transaction.bank_serial_no or "",
                decimal_cell(transaction.matched_amount),
            ]
        )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column_cells in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 12), 36)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"ledger-detail-{store_id}-{period}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
