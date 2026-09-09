import io
import json
from datetime import datetime
from decimal import Decimal
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import ApprovalInstance, ApprovalTemplate, ExpenseCategory, ExpenseItem, TemplateFieldMapping
from app.modules.matching.router import candidate_score


def create_expense_category_pair(session: Session, category_l1: str, category_l2: str) -> None:
    parent = ExpenseCategory(name=category_l1, parent_id=None, sort_order=1)
    session.add(parent)
    session.flush()
    session.add(ExpenseCategory(name=category_l2, parent_id=parent.id, sort_order=1))
    session.flush()


def test_store_ledger_and_close_flow(client: TestClient) -> None:
    store_response = client.post("/api/stores", json={"name": "蘑说测试店"})
    assert store_response.status_code == 201
    store_id = store_response.json()["data"]["id"]

    ledger_response = client.post(
        "/api/ledgers",
        json={"store_id": store_id, "period": "2026-08"},
    )
    assert ledger_response.status_code == 201
    ledger_id = ledger_response.json()["data"]["id"]

    close_response = client.post(f"/api/ledgers/{ledger_id}/close", json={"operator": "tester"})
    assert close_response.status_code == 200
    assert close_response.json()["data"]["status"] == "closed"

    blocked_response = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "封账后支出",
            "amount": "20.00",
        },
    )
    assert blocked_response.status_code == 409


def test_close_ledger_requires_clean_close_check(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说封账检查店"}).json()["data"]["id"]
    ledger_id = client.post(
        "/api/ledgers",
        json={"store_id": store_id, "period": "2026-08"},
    ).json()["data"]["id"]
    client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "待付款支出",
            "amount": "100.00",
        },
    )

    check_response = client.get(f"/api/ledgers/{ledger_id}/close-check")
    assert check_response.status_code == 200
    check_data = check_response.json()["data"]
    assert check_data["can_close"] is False
    assert check_data["unpaid_expense_count"] == 1

    close_response = client.post(f"/api/ledgers/{ledger_id}/close", json={"operator": "tester"})
    assert close_response.status_code == 409


def test_update_store_status_and_audit_log(client: TestClient) -> None:
    store_response = client.post("/api/stores", json={"name": "蘑说门店维护店"})
    assert store_response.status_code == 201
    store_id = store_response.json()["data"]["id"]

    update_response = client.patch(
        f"/api/stores/{store_id}",
        json={"contact_person": "门店负责人", "status": "inactive"},
    )
    assert update_response.status_code == 200
    data = update_response.json()["data"]
    assert data["contact_person"] == "门店负责人"
    assert data["status"] == "inactive"

    logs = client.get("/api/audit-logs?resource_type=store&page_size=20").json()["data"]["items"]
    assert any(log["action"] == "store.update" and log["resource_id"] == store_id for log in logs)


def test_confirm_match_updates_payment_status(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说匹配店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})

    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "水电费",
            "amount": "1280.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "1280.00",
            "counterparty_name": "供电公司",
        },
    ).json()["data"]["id"]

    match_id = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": bank_id,
            "amount": "1280.00",
            "confidence": "92.00",
            "reason": "金额一致",
        },
    ).json()["data"]["id"]

    confirm_response = client.post(f"/api/matches/{match_id}/confirm?operator=tester")
    assert confirm_response.status_code == 200
    assert confirm_response.json()["data"]["status"] == "confirmed"

    expense_items = client.get("/api/expense-items").json()["data"]["items"]
    assert expense_items[0]["payment_status"] == "paid"

    reject_confirmed_response = client.post(f"/api/matches/{match_id}/reject")
    assert reject_confirmed_response.status_code == 409


def test_confirm_match_keeps_bank_transaction_store_and_period(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说未归属流水匹配店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-20",
            "description": "门店报销付款",
            "amount": "300.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "300.00",
            "summary": "门店报销付款",
        },
    ).json()["data"]["id"]

    candidates = client.get(f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}").json()["data"]
    assert candidates["remaining_amount"] == "300.00"
    assert candidates["candidates"][0]["expense_item"]["id"] == expense_id

    match_id = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": bank_id,
            "amount": "300.00",
            "accounting_period": "2026-08",
            "category_l1": "日常支出",
            "category_l2": "门店零星报销",
        },
    ).json()["data"]["id"]
    confirm_response = client.post(f"/api/matches/{match_id}/confirm?operator=tester")
    assert confirm_response.status_code == 200

    bank_transaction = client.get("/api/bank-transactions?direction=expense").json()["data"]["items"][0]
    assert bank_transaction["store_id"] == store_id
    assert bank_transaction["ledger_period"] == "2026-08"
    assert bank_transaction["matched_amount"] == "300.00"

    records = client.get("/api/matches/reconciliation/records?accounting_period=2026-08").json()["data"]
    assert records["total"] == 1
    assert records["items"][0]["match"]["accounting_period"] == "2026-08"
    assert records["items"][0]["match"]["bank_occurred"] is True
    assert records["items"][0]["expense_item"]["id"] == expense_id
    assert records["items"][0]["expense_item"]["category_l1"] == "日常支出"
    assert records["items"][0]["expense_item"]["category_l2"] == "门店零星报销"

    update_response = client.patch(
        f"/api/matches/reconciliation/records/{match_id}",
        json={
            "accounting_period": "2026-10",
            "bank_occurred": False,
            "category_l1": "人工成本",
            "category_l2": "员工工资",
            "reason": "银行暂未实际发生",
        },
    )
    assert update_response.status_code == 200
    assert update_response.json()["data"]["accounting_period"] == "2026-10"
    assert update_response.json()["data"]["bank_occurred"] is False

    updated_records = client.get("/api/matches/reconciliation/records?accounting_period=2026-10").json()["data"]
    assert updated_records["total"] == 1
    assert updated_records["items"][0]["match"]["bank_occurred"] is False
    assert updated_records["items"][0]["expense_item"]["category_l1"] == "人工成本"
    assert updated_records["items"][0]["expense_item"]["category_l2"] == "员工工资"

    unmatch_response = client.post(f"/api/matches/reconciliation/records/{match_id}/unmatch")
    assert unmatch_response.status_code == 200
    assert unmatch_response.json()["data"]["status"] == "rejected"

    unmatched_records = client.get("/api/matches/reconciliation/records?accounting_period=2026-10").json()["data"]
    assert unmatched_records["total"] == 0
    bank_transaction = client.get("/api/bank-transactions?direction=expense").json()["data"]["items"][0]
    assert bank_transaction["matched_amount"] == "0.00"
    assert bank_transaction["store_id"] == store_id
    assert bank_transaction["ledger_period"] == "2026-08"
    expense_items = client.get("/api/expense-items?payment_status=unpaid").json()["data"]["items"]
    assert expense_items[0]["id"] == expense_id


def test_reconciliation_records_include_cross_period_match(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说跨账期记录店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-09"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-31",
            "description": "跨账期支出",
            "amount": "100.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-09",
            "occurred_at": "2026-09-01T10:00:00",
            "direction": "expense",
            "amount": "100.00",
        },
    ).json()["data"]["id"]
    match_id = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": bank_id,
            "amount": "100.00",
            "accounting_period": "2026-08",
        },
    ).json()["data"]["id"]
    assert client.post(f"/api/matches/{match_id}/confirm").status_code == 200

    records = client.get(
        f"/api/matches/reconciliation/records?store_id={store_id}&accounting_period=2026-09"
    ).json()["data"]
    assert records["total"] == 1
    assert records["items"][0]["match"]["accounting_period"] == "2026-08"


def test_candidate_score_only_uses_bank_and_approval_amounts() -> None:
    bank_transaction = SimpleNamespace(amount=Decimal("10000.00"), matched_amount=Decimal("0.00"))
    expense = SimpleNamespace(amount=Decimal("10000.00"))

    exact_score, exact_reason = candidate_score(
        bank_transaction,
        expense,
        Decimal("10000.00"),
        approval_total_amount=Decimal("10000.00"),
    )
    close_score, close_reason = candidate_score(
        bank_transaction,
        expense,
        Decimal("9000.00"),
        approval_total_amount=Decimal("9000.00"),
    )

    assert exact_score == Decimal("100.00")
    assert close_score == Decimal("90.00")
    assert "门店" not in exact_reason and "日期" not in exact_reason
    assert "门店" not in close_reason and "日期" not in close_reason


def test_create_match_candidate_is_idempotent_for_existing_candidate(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说重复候选店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-20",
            "description": "重复候选支出",
            "amount": "1980.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "1980.00",
            "summary": "重复候选流水",
        },
    ).json()["data"]["id"]
    payload = {
        "expense_item_id": expense_id,
        "bank_transaction_id": bank_id,
        "amount": "1980.00",
        "accounting_period": "2026-08",
        "bank_occurred": True,
    }

    first = client.post("/api/matches", json=payload)
    second = client.post("/api/matches", json={**payload, "reason": "重复提交"})

    assert first.status_code == 201
    assert second.status_code == 201
    assert second.json()["data"]["id"] == first.json()["data"]["id"]
    assert second.json()["data"]["reason"] == "重复提交"


def test_create_match_candidate_rejects_existing_confirmed_match(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说重复确认店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-20",
            "description": "重复确认支出",
            "amount": "1980.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "1980.00",
            "summary": "重复确认流水",
        },
    ).json()["data"]["id"]
    payload = {
        "expense_item_id": expense_id,
        "bank_transaction_id": bank_id,
        "amount": "1980.00",
        "accounting_period": "2026-08",
        "bank_occurred": True,
    }
    match_id = client.post("/api/matches", json=payload).json()["data"]["id"]
    client.post(f"/api/matches/{match_id}/confirm?operator=tester")

    duplicate = client.post("/api/matches", json=payload)

    assert duplicate.status_code == 409
    assert duplicate.json()["detail"] == "This bank transaction is already matched to the approval"


def test_create_match_candidate_reuses_rejected_existing_match(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说重新匹配店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-20",
            "description": "重新匹配支出",
            "amount": "1980.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "1980.00",
            "summary": "重新匹配流水",
        },
    ).json()["data"]["id"]
    payload = {
        "expense_item_id": expense_id,
        "bank_transaction_id": bank_id,
        "amount": "1980.00",
        "accounting_period": "2026-08",
        "bank_occurred": True,
        "reason": "首次匹配",
    }
    first = client.post("/api/matches", json=payload)
    assert first.status_code == 201
    match_id = first.json()["data"]["id"]
    confirm = client.post(f"/api/matches/{match_id}/confirm?operator=tester")
    assert confirm.status_code == 200
    unmatch = client.post(f"/api/matches/reconciliation/records/{match_id}/unmatch")
    assert unmatch.status_code == 200
    assert unmatch.json()["data"]["status"] == "rejected"

    retry = client.post("/api/matches", json={**payload, "reason": "重新匹配"})

    assert retry.status_code == 201
    assert retry.json()["data"]["id"] == match_id
    assert retry.json()["data"]["status"] == "candidate"
    assert retry.json()["data"]["reason"] == "重新匹配"


def test_bank_transaction_cannot_match_multiple_expense_items(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说流水拆分匹配店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    first_expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-20",
            "description": "第一个审批单",
            "amount": "400.00",
        },
    ).json()["data"]["id"]
    second_expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-21",
            "description": "第二个审批单",
            "amount": "600.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-22T10:30:00",
            "direction": "expense",
            "amount": "1000.00",
            "summary": "一笔流水合并支付两个审批",
        },
    ).json()["data"]["id"]

    first = client.post(
        "/api/matches",
        json={
            "expense_item_id": first_expense_id,
            "bank_transaction_id": bank_id,
            "amount": "400.00",
            "accounting_period": "2026-08",
        },
    )
    second = client.post(
        "/api/matches",
        json={
            "expense_item_id": second_expense_id,
            "bank_transaction_id": bank_id,
            "amount": "600.00",
            "accounting_period": "2026-08",
        },
    )

    assert first.status_code == 201
    assert second.status_code == 409
    assert second.json()["detail"] == "This bank transaction is already matched to an approval"

    first_confirm = client.post(f"/api/matches/{first.json()['data']['id']}/confirm?operator=tester")
    assert first_confirm.status_code == 200

    bank_transaction = client.get(f"/api/bank-transactions?store_id={store_id}&page_size=20").json()["data"]["items"][0]
    assert bank_transaction["matched_amount"] == "400.00"

    third_expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-22",
            "description": "第三个审批单",
            "amount": "1.00",
        },
    ).json()["data"]["id"]
    over_match = client.post(
        "/api/matches",
        json={
            "expense_item_id": third_expense_id,
            "bank_transaction_id": bank_id,
            "amount": "1.00",
            "accounting_period": "2026-08",
        },
    )
    assert over_match.status_code == 409
    assert over_match.json()["detail"] == "This bank transaction is already matched to an approval"


def test_reconciliation_candidate_search_returns_parsed_approval_lines_without_backfill(
    client: TestClient,
    session: Session,
) -> None:
    store_id = client.post("/api/stores", json={"name": "菌山集阳江新达城店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    create_expense_category_pair(session, "门店零星报销", "零食物料")
    template = ApprovalTemplate(process_code="PROC-RECON", name="门店支出报销")
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-instance-for-search",
        approval_no="202608242148000482025",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 24, 21, 48, 49),
        raw_payload=json.dumps(
            {
                "title": "林燕玲提交的门店支出报销",
                "business_id": "202608242148000482025",
                "create_time": "2026-08-24 21:48:49",
                "form_component_values": [
                    {"name": "报销日期", "value": "2026-08-24"},
                    {"name": "支出类型", "value": "门店零星报销"},
                    {"name": "汇总金额（元）", "value": "4020.04"},
                    {"name": "收款账户", "value": "林燕玲"},
                ],
            },
            ensure_ascii=False,
        ),
    )
    session.add(approval)
    session.flush()
    session.add_all(
        [
            ExpenseItem(
                store_id=store_id,
                ledger_period="2026-08",
                expense_date=datetime(2026, 8, 24),
                description="爆米花",
                amount="104.50",
                source="dingtalk",
                source_document_id="approval-instance-for-search:1",
            ),
            ExpenseItem(
                store_id=store_id,
                ledger_period="2026-08",
                expense_date=datetime(2026, 8, 24),
                description="薄荷糖",
                amount="27.79",
                source="dingtalk",
                source_document_id="approval-instance-for-search:2",
            ),
        ]
    )
    session.commit()
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-25T10:30:00",
            "direction": "expense",
            "amount": "4020.04",
            "summary": "报销付款 202608242148000482025",
        },
    ).json()["data"]["id"]

    candidates = client.get(f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}").json()["data"][
        "candidates"
    ]

    assert len(candidates) == 1
    assert {candidate["approval_instance"]["approval_no"] for candidate in candidates} == {"202608242148000482025"}
    assert candidates[0]["expense_item"]["description"] in {"爆米花", "薄荷糖"}
    assert candidates[0]["remaining_amount"] in {"104.50", "27.79"}
    assert all(":" in candidate["expense_item"]["source_document_id"] for candidate in candidates)
    assert (
        session.query(ExpenseItem)
        .filter(ExpenseItem.source_document_id == "approval-instance-for-search")
        .count()
        == 0
    )

    selected_candidate = candidates[0]
    match_response = client.post(
        "/api/matches",
        json={
            "expense_item_id": selected_candidate["expense_item"]["id"],
            "bank_transaction_id": bank_id,
            "amount": selected_candidate["remaining_amount"],
            "accounting_period": "2026-08",
            "bank_occurred": True,
            "category_l1": "门店零星报销",
            "category_l2": "零食物料",
        },
    )
    assert match_response.status_code == 201
    matched_expense = session.get(ExpenseItem, selected_candidate["expense_item"]["id"])
    assert matched_expense is not None
    assert matched_expense.category_l1 == "门店零星报销"
    assert matched_expense.category_l2 == "零食物料"
    assert json.loads(matched_expense.user_edited_fields_json) == ["category_l1", "category_l2"]

    searched_candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}&approval_no=202608242148000482025"
    ).json()["data"]["candidates"]
    assert len(searched_candidates) == 1
    assert searched_candidates[0]["expense_item"]["description"] in {"爆米花", "薄荷糖"}


def test_reconciliation_candidate_search_skips_zero_amount_approval_lines(
    client: TestClient,
    session: Session,
) -> None:
    store_id = client.post("/api/stores", json={"name": "零金额明细候选店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    template = ApprovalTemplate(process_code="PROC-ZERO-LINE", name="门店支出报销")
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-with-zero-line",
        approval_no="202609030332000374852",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 9, 3, 3, 32, 37),
    )
    session.add(approval)
    session.flush()
    session.add_all(
        [
            ExpenseItem(
                store_id=store_id,
                ledger_period="2026-08",
                expense_date=datetime(2026, 9, 3),
                description="有效明细",
                amount="44540.20",
                source="dingtalk",
                source_document_id="approval-with-zero-line:line-1",
            ),
            ExpenseItem(
                store_id=store_id,
                ledger_period="2026-08",
                expense_date=datetime(2026, 9, 3),
                description="零金额明细",
                amount="0.00",
                source="dingtalk",
                source_document_id="approval-with-zero-line:line-2",
            ),
        ]
    )
    session.commit()
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-09",
            "occurred_at": "2026-08-21T00:00:00",
            "direction": "expense",
            "amount": "50000.00",
            "summary": "零金额明细候选",
        },
    ).json()["data"]["id"]

    response = client.get(
        f"/api/matches/reconciliation/candidates?store_id={store_id}"
        f"&approval_only=true&page_size=100&bank_transaction_id={bank_id}"
        "&approval_no=202609030332000374852"
    )

    assert response.status_code == 200
    candidates = response.json()["data"]["candidates"]
    assert len(candidates) == 1
    assert candidates[0]["expense_item"]["description"] == "有效明细"
    assert candidates[0]["expense_item"]["amount"] == "44540.20"


def test_reconciliation_candidates_can_list_store_approvals_without_bank_transaction(
    client: TestClient,
    session: Session,
) -> None:
    store_id = client.post("/api/stores", json={"name": "无流水候选门店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    template = ApprovalTemplate(process_code="PROC-NO-BANK-RECON", name="门店支出报销")
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-no-bank-candidate",
        approval_no="202608319900",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 31, 10, 0, 0),
        raw_payload=json.dumps(
            {
                "title": "李四提交的门店支出报销",
                "business_id": "202608319900",
                "create_time": "2026-08-31 10:00:00",
                "form_component_values": [
                    {"name": "报销日期", "value": "2026-08-31"},
                    {"name": "汇总金额（元）", "value": "388.00"},
                ],
            },
            ensure_ascii=False,
        ),
    )
    session.add(approval)
    session.flush()
    session.add(
        ExpenseItem(
            store_id=store_id,
            ledger_period="2026-08",
            expense_date=datetime(2026, 8, 31),
            description="李四提交的门店支出报销",
            amount="388.00",
            source="dingtalk",
            source_document_id="approval-no-bank-candidate",
            approval_instance_id=approval.id,
            approval_line_no=1,
            approval_line_key="whole-approval",
            approval_line_source_type="whole_approval",
        )
    )
    session.commit()

    response = client.get(f"/api/matches/reconciliation/candidates?store_id={store_id}&approval_only=true")

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["bank_transaction"] is None
    assert data["remaining_amount"] == "0.00"
    assert len(data["candidates"]) == 1
    assert data["candidates"][0]["approval_instance"]["approval_no"] == "202608319900"
    assert data["candidates"][0]["expense_item"]["amount"] == "388.00"


def test_reconciliation_candidate_backfill_uses_template_business_amount_mapping(
    client: TestClient,
    session: Session,
) -> None:
    store_id = client.post("/api/stores", json={"name": "筹建报销测试店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    template = ApprovalTemplate(process_code="PROC-BUILD-RECON", name="门店筹建报销")
    session.add(template)
    session.flush()
    session.add(
        TemplateFieldMapping(
            template_id=template.id,
            standard_field="amount",
            display_label="单据总金额",
            source_field_id="build-amount-field",
            source_field_name="金额（元）",
            source_path="field:build-amount-field",
            field_type="MoneyField",
            is_required=True,
        )
    )
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="build-approval-instance",
        approval_no="202608300088",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 30, 10, 0, 0),
        raw_payload=json.dumps(
            {
                "title": "张三提交的门店筹建报销",
                "business_id": "202608300088",
                "create_time": "2026-08-30 10:00:00",
                "form_component_values": [
                    {"name": "申请日期", "value": "2026-08-30"},
                    {"id": "build-amount-field", "name": "金额（元）", "value": "1288.66"},
                ],
            },
            ensure_ascii=False,
        ),
    )
    session.add(approval)
    session.flush()
    session.add(
        ExpenseItem(
            store_id=store_id,
            ledger_period="2026-08",
            expense_date=datetime(2026, 8, 30),
            description="张三提交的门店筹建报销",
            amount="1288.66",
            source="dingtalk",
            source_document_id="build-approval-instance",
            approval_instance_id=approval.id,
            approval_line_no=1,
            approval_line_key="whole-approval",
            approval_line_source_type="whole_approval",
        )
    )
    session.commit()
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-30T11:30:00",
            "direction": "expense",
            "amount": "1288.66",
            "summary": "筹建报销 202608300088",
        },
    ).json()["data"]["id"]

    candidates = client.get(f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}").json()["data"][
        "candidates"
    ]

    assert len(candidates) == 1
    assert candidates[0]["approval_instance"]["approval_no"] == "202608300088"
    assert candidates[0]["remaining_amount"] == "1288.66"


def test_reconciliation_candidates_ignore_disabled_approval_templates(
    client: TestClient,
    session: Session,
) -> None:
    store_id = client.post("/api/stores", json={"name": "停用模板候选门店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    template = ApprovalTemplate(process_code="PROC-DISABLED-RECON", name="停用对账模板", is_enabled=False)
    session.add(template)
    session.flush()
    session.add(
        ApprovalInstance(
            template_id=template.id,
            dingtalk_instance_id="disabled-approval-instance",
            approval_no="202608300001",
            store_id=store_id,
            approval_status="agree",
            submit_at=datetime(2026, 8, 30, 10, 0, 0),
            raw_payload=json.dumps(
                {
                    "title": "停用模板审批",
                    "business_id": "202608300001",
                    "form_component_values": [
                        {"name": "报销日期", "value": "2026-08-30"},
                        {"name": "汇总金额（元）", "value": "600.00"},
                    ],
                },
                ensure_ascii=False,
            ),
        )
    )
    expense = ExpenseItem(
        store_id=store_id,
        ledger_period="2026-08",
        expense_date=datetime(2026, 8, 30),
        description="停用模板审批",
        amount="600.00",
        source="dingtalk",
        source_document_id="disabled-approval-instance",
    )
    session.add(expense)
    session.commit()
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-30T11:00:00",
            "direction": "expense",
            "amount": "600.00",
            "summary": "停用模板审批",
        },
    ).json()["data"]["id"]

    candidates = client.get(f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}").json()["data"][
        "candidates"
    ]

    assert candidates == []


def test_reconciliation_candidates_are_limited_by_bank_store(client: TestClient) -> None:
    bank_store_id = client.post("/api/stores", json={"name": "流水归属店"}).json()["data"]["id"]
    approval_store_id = client.post("/api/stores", json={"name": "审批归属店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": bank_store_id, "period": "2026-08"})
    client.post("/api/ledgers", json={"store_id": approval_store_id, "period": "2026-08"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": approval_store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-20",
            "description": "跨门店候选审批单",
            "amount": "888.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": bank_store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "888.00",
            "summary": "跨门店候选测试",
        },
    ).json()["data"]["id"]

    candidates = client.get(f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}").json()["data"][
        "candidates"
    ]
    assert all(candidate["expense_item"]["store_id"] == bank_store_id for candidate in candidates)
    assert all(candidate["expense_item"]["id"] != expense_id for candidate in candidates)

    filtered_candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}&store_id={bank_store_id}"
    ).json()["data"]["candidates"]
    assert all(candidate["expense_item"]["store_id"] == bank_store_id for candidate in filtered_candidates)

    mismatch_response = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": bank_id,
            "amount": "888.00",
        },
    )
    assert mismatch_response.status_code == 409
    assert mismatch_response.json()["detail"] == "Store mismatch"


def test_reconciliation_candidates_can_filter_by_ledger_period(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "账期候选过滤店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-09"})
    august_expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-20",
            "description": "八月审批费用",
            "amount": "500.00",
        },
    ).json()["data"]["id"]
    september_expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-09",
            "expense_date": "2026-09-20",
            "description": "九月审批费用",
            "amount": "500.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "500.00",
            "summary": "账期候选过滤",
        },
    ).json()["data"]["id"]

    candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}&store_id={store_id}&ledger_period=2026-08"
    ).json()["data"]["candidates"]
    candidate_ids = {candidate["expense_item"]["id"] for candidate in candidates}

    assert august_expense_id in candidate_ids
    assert september_expense_id not in candidate_ids


def test_reconciliation_candidates_can_filter_to_approval_documents_only(
    client: TestClient,
    session: Session,
) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说审批候选店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    manual_expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "expense_date": "2026-08-20",
            "description": "手工支出",
            "amount": "500.00",
        },
    ).json()["data"]["id"]
    template = ApprovalTemplate(process_code="PROC-APPROVAL-ONLY", name="门店支出报销", is_enabled=True)
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-only-instance",
        approval_no="202608300188",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 20, 10, 0, 0),
        raw_payload=json.dumps(
            {
                "title": "审批候选测试",
                "business_id": "202608300188",
                "form_component_values": [{"name": "汇总金额（元）", "value": "500.00"}],
            },
            ensure_ascii=False,
        ),
    )
    session.add(approval)
    session.flush()
    session.add(
        ExpenseItem(
            store_id=store_id,
            ledger_period="2026-08",
            expense_date=datetime(2026, 8, 20),
            description="审批候选测试",
            amount="500.00",
            source="dingtalk",
            source_document_id="approval-only-instance",
            approval_instance_id=approval.id,
            approval_line_no=1,
            approval_line_key="whole-approval",
            approval_line_source_type="whole_approval",
        )
    )
    sample_template = ApprovalTemplate(process_code="PROC-SAMPLE-APPROVAL-ONLY", name="样例审批", is_enabled=True)
    session.add(sample_template)
    session.flush()
    session.add(
        ApprovalInstance(
            template_id=sample_template.id,
            dingtalk_instance_id="sample-approval-only-instance",
            approval_no="SAMPLE-审批样例",
            store_id=store_id,
            approval_status="agree",
            submit_at=datetime(2026, 8, 20, 10, 0, 0),
            raw_payload=json.dumps(
                {
                    "title": "样例审批",
                    "business_id": "SAMPLE-审批样例",
                    "form_component_values": [{"name": "汇总金额（元）", "value": "500.00"}],
                },
                ensure_ascii=False,
            ),
        )
    )
    session.commit()
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "500.00",
            "summary": "审批候选测试",
        },
    ).json()["data"]["id"]

    all_candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}&store_id={store_id}"
    ).json()["data"]["candidates"]
    approval_candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={bank_id}&store_id={store_id}&approval_only=true"
    ).json()["data"]["candidates"]

    assert any(candidate["expense_item"]["id"] == manual_expense_id for candidate in all_candidates)
    assert [candidate["approval_instance"]["approval_no"] for candidate in approval_candidates] == ["202608300188"]


def test_reconciliation_candidates_exclude_expense_with_active_match(
    client: TestClient,
    session: Session,
) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说候选占用店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    create_expense_category_pair(session, "门店支出", "零星报销")
    template = ApprovalTemplate(process_code="PROC-ACTIVE-MATCH", name="门店支出报销", is_enabled=True)
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-active-match",
        approval_no="202608300299",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 20, 10, 0, 0),
        raw_payload=json.dumps(
            {
                "title": "候选占用审批",
                "business_id": "202608300299",
                "form_component_values": [{"name": "汇总金额（元）", "value": "500.00"}],
            },
            ensure_ascii=False,
        ),
    )
    session.add(approval)
    session.flush()
    session.add(
        ExpenseItem(
            store_id=store_id,
            ledger_period="2026-08",
            expense_date=datetime(2026, 8, 20),
            description="候选占用审批",
            amount="500.00",
            source="dingtalk",
            source_document_id="approval-active-match",
            approval_instance_id=approval.id,
            approval_line_no=1,
            approval_line_key="whole-approval",
            approval_line_source_type="whole_approval",
        )
    )
    session.commit()
    first_bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "500.00",
            "summary": "第一条流水",
        },
    ).json()["data"]["id"]
    second_bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-22T10:30:00",
            "direction": "expense",
            "amount": "500.00",
            "summary": "第二条流水",
        },
    ).json()["data"]["id"]
    first_candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={first_bank_id}&store_id={store_id}&approval_only=true"
    ).json()["data"]["candidates"]
    expense_id = first_candidates[0]["expense_item"]["id"]
    create_response = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": first_bank_id,
            "amount": "500.00",
            "accounting_period": "2026-08",
            "category_l1": "门店支出",
            "category_l2": "零星报销",
        },
    )
    assert create_response.status_code == 201
    match_id = create_response.json()["data"]["id"]

    confirm_response = client.post(f"/api/matches/{match_id}/confirm?operator=tester")
    assert confirm_response.status_code == 200

    second_candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={second_bank_id}&store_id={store_id}&approval_only=true"
    ).json()["data"]["candidates"]

    assert all(candidate["expense_item"]["id"] != expense_id for candidate in second_candidates)

    searched_candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={second_bank_id}&store_id={store_id}&approval_only=true&approval_no=202608300299"
    ).json()["data"]["candidates"]
    assert len(searched_candidates) == 1
    assert searched_candidates[0]["expense_item"]["id"] == expense_id
    assert searched_candidates[0]["remaining_amount"] == "0.00"
    duplicate_response = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": second_bank_id,
            "amount": "500.00",
            "accounting_period": "2026-08",
            "category_l1": "门店支出",
            "category_l2": "零星报销",
        },
    )
    assert duplicate_response.status_code == 201


def test_reconciliation_search_allows_partial_approval_remaining_match(
    client: TestClient,
    session: Session,
) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说审批部分付款店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    create_expense_category_pair(session, "门店支出", "零星报销")
    template = ApprovalTemplate(process_code="PROC-PARTIAL-APPROVAL", name="门店支出报销", is_enabled=True)
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-partial-match",
        approval_no="202608300588",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 20, 10, 0, 0),
    )
    session.add(approval)
    session.flush()
    expense = ExpenseItem(
        store_id=store_id,
        ledger_period="2026-08",
        expense_date=datetime(2026, 8, 20),
        description="审批分两笔付款",
        amount="1000.00",
        source="dingtalk",
        source_document_id="approval-partial-match",
        approval_instance_id=approval.id,
        approval_line_no=1,
        approval_line_key="whole-approval",
        approval_line_source_type="whole_approval",
    )
    session.add(expense)
    session.commit()
    first_bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "400.00",
            "summary": "审批第一笔付款",
        },
    ).json()["data"]["id"]
    second_bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-22T10:30:00",
            "direction": "expense",
            "amount": "600.00",
            "summary": "审批第二笔付款",
        },
    ).json()["data"]["id"]
    first_match = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense.id,
            "bank_transaction_id": first_bank_id,
            "amount": "400.00",
            "accounting_period": "2026-08",
            "category_l1": "门店支出",
            "category_l2": "零星报销",
        },
    )
    assert first_match.status_code == 201
    assert client.post(f"/api/matches/{first_match.json()['data']['id']}/confirm?operator=tester").status_code == 200

    hidden_candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={second_bank_id}&store_id={store_id}&approval_only=true"
    ).json()["data"]["candidates"]
    assert all(candidate["expense_item"]["id"] != expense.id for candidate in hidden_candidates)

    searched_candidates = client.get(
        f"/api/matches/reconciliation/candidates?bank_transaction_id={second_bank_id}"
        f"&store_id={store_id}&approval_only=true&approval_no=202608300588"
    ).json()["data"]["candidates"]
    assert len(searched_candidates) == 1
    assert searched_candidates[0]["expense_item"]["id"] == expense.id
    assert searched_candidates[0]["remaining_amount"] == "600.00"

    second_match = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense.id,
            "bank_transaction_id": second_bank_id,
            "amount": "600.00",
            "accounting_period": "2026-08",
            "category_l1": "门店支出",
            "category_l2": "零星报销",
        },
    )
    assert second_match.status_code == 201


def test_approval_match_requires_category_before_save(client: TestClient, session: Session) -> None:
    store_id = client.post("/api/stores", json={"name": "审批匹配分类必填店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    create_expense_category_pair(session, "房租水电", "水电费")
    template = ApprovalTemplate(process_code="PROC-CATEGORY-REQUIRED", name="门店支出报销", is_enabled=True)
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-category-required",
        approval_no="202608300399",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 20, 10, 0, 0),
    )
    session.add(approval)
    session.flush()
    expense = ExpenseItem(
        store_id=store_id,
        ledger_period="2026-08",
        expense_date=datetime(2026, 8, 20),
        description="未分类审批支出",
        amount="500.00",
        source="dingtalk",
        source_document_id="approval-category-required",
        approval_instance_id=approval.id,
    )
    session.add(expense)
    session.commit()
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "500.00",
            "summary": "审批匹配分类必填",
        },
    ).json()["data"]["id"]
    payload = {
        "expense_item_id": expense.id,
        "bank_transaction_id": bank_id,
        "amount": "500.00",
        "accounting_period": "2026-08",
    }

    missing_category = client.post("/api/matches", json=payload)
    invalid_category = client.post(
        "/api/matches",
        json={**payload, "category_l1": "门店支出", "category_l2": "零星报销"},
    )
    with_category = client.post(
        "/api/matches",
        json={**payload, "category_l1": "房租水电", "category_l2": "水电费"},
    )

    assert missing_category.status_code == 409
    assert missing_category.json()["detail"] == "请选择费用分类后再确认匹配"
    assert invalid_category.status_code == 409
    assert invalid_category.json()["detail"] == "请选择费用分类后再确认匹配"
    assert with_category.status_code == 201
    confirm_response = client.post(f"/api/matches/{with_category.json()['data']['id']}/confirm?operator=tester")
    assert confirm_response.status_code == 200


def test_approval_match_rejects_revenue_fee_category(client: TestClient, session: Session) -> None:
    store_id = client.post("/api/stores", json={"name": "审批匹配手续费禁用店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    create_expense_category_pair(session, "手续费", "美团手续费")
    template = ApprovalTemplate(process_code="PROC-FEE-CATEGORY-REJECT", name="门店支出报销", is_enabled=True)
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-fee-category-reject",
        approval_no="202608300499",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 20, 10, 0, 0),
    )
    session.add(approval)
    session.flush()
    expense = ExpenseItem(
        store_id=store_id,
        ledger_period="2026-08",
        expense_date=datetime(2026, 8, 20),
        description="误选手续费分类审批支出",
        amount="500.00",
        source="dingtalk",
        source_document_id="approval-fee-category-reject",
        approval_instance_id=approval.id,
    )
    session.add(expense)
    session.commit()
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "500.00",
            "summary": "审批匹配不能选手续费",
        },
    ).json()["data"]["id"]

    response = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense.id,
            "bank_transaction_id": bank_id,
            "amount": "500.00",
            "accounting_period": "2026-08",
            "category_l1": "手续费",
            "category_l2": "美团手续费",
        },
    )

    assert response.status_code == 409
    assert response.json()["detail"] == "请选择费用分类后再确认匹配"


def test_approval_match_ignores_synced_category_until_user_saves(client: TestClient, session: Session) -> None:
    store_id = client.post("/api/stores", json={"name": "审批同步分类不回显店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    create_expense_category_pair(session, "房租水电", "水电费")
    template = ApprovalTemplate(process_code="PROC-SYNCED-CATEGORY", name="门店支出报销", is_enabled=True)
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-synced-category",
        approval_no="202608300498",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 20, 10, 0, 0),
    )
    session.add(approval)
    session.flush()
    expense = ExpenseItem(
        store_id=store_id,
        ledger_period="2026-08",
        expense_date=datetime(2026, 8, 20),
        description="同步已有分类审批支出",
        amount="500.00",
        category_l1="房租水电",
        category_l2="水电费",
        source="dingtalk",
        source_document_id="approval-synced-category",
        approval_instance_id=approval.id,
    )
    session.add(expense)
    session.commit()
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "500.00",
            "summary": "审批同步分类不回显",
        },
    ).json()["data"]["id"]

    payload = {
        "expense_item_id": expense.id,
        "bank_transaction_id": bank_id,
        "amount": "500.00",
        "accounting_period": "2026-08",
    }

    synced_category = client.post("/api/matches", json=payload)
    user_selected_category = client.post(
        "/api/matches",
        json={**payload, "category_l1": "房租水电", "category_l2": "水电费"},
    )

    assert synced_category.status_code == 409
    assert synced_category.json()["detail"] == "请选择费用分类后再确认匹配"
    assert user_selected_category.status_code == 201
    session.refresh(expense)
    assert json.loads(expense.user_edited_fields_json) == ["category_l1", "category_l2"]


def test_approval_match_accepts_single_level_category_payload(client: TestClient, session: Session) -> None:
    store_id = client.post("/api/stores", json={"name": "审批单级分类店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    session.add(ExpenseCategory(name="房租水电", parent_id=None, sort_order=1))
    template = ApprovalTemplate(process_code="PROC-SINGLE-CATEGORY", name="门店支出报销", is_enabled=True)
    session.add(template)
    session.flush()
    approval = ApprovalInstance(
        template_id=template.id,
        dingtalk_instance_id="approval-single-category",
        approval_no="202608300598",
        store_id=store_id,
        approval_status="agree",
        submit_at=datetime(2026, 8, 20, 10, 0, 0),
    )
    session.add(approval)
    session.flush()
    expense = ExpenseItem(
        store_id=store_id,
        ledger_period="2026-08",
        expense_date=datetime(2026, 8, 20),
        description="单级分类审批支出",
        amount="300.00",
        source="dingtalk",
        source_document_id="approval-single-category",
        approval_instance_id=approval.id,
    )
    session.add(expense)
    session.commit()
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-21T10:30:00",
            "direction": "expense",
            "amount": "300.00",
            "summary": "审批单级分类",
        },
    ).json()["data"]["id"]

    response = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense.id,
            "bank_transaction_id": bank_id,
            "amount": "300.00",
            "accounting_period": "2026-08",
            "category_l1": "房租水电",
            "category_l2": "房租水电",
        },
    )

    assert response.status_code == 201
    session.refresh(expense)
    assert expense.category_l1 == "房租水电"
    assert expense.category_l2 is None


def test_confirm_multiple_matches_updates_partial_and_paid_status(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说多笔付款店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})

    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "分两笔支付的货款",
            "amount": "1000.00",
        },
    ).json()["data"]["id"]
    first_bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "400.00",
            "counterparty_name": "供应商A",
        },
    ).json()["data"]["id"]
    second_bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-14T10:30:00",
            "direction": "expense",
            "amount": "600.00",
            "counterparty_name": "供应商A",
        },
    ).json()["data"]["id"]

    first_match_id = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": first_bank_id,
            "amount": "400.00",
            "reason": "第一笔付款",
        },
    ).json()["data"]["id"]
    first_confirm_response = client.post(f"/api/matches/{first_match_id}/confirm?operator=tester")
    assert first_confirm_response.status_code == 200

    expense_items = client.get("/api/expense-items?payment_status=unpaid,partial_paid").json()["data"]["items"]
    assert expense_items[0]["id"] == expense_id
    assert expense_items[0]["payment_status"] == "partial_paid"

    over_match_response = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": second_bank_id,
            "amount": "601.00",
            "reason": "超额付款",
        },
    )
    assert over_match_response.status_code == 201
    over_match_id = over_match_response.json()["data"]["id"]
    assert client.post(f"/api/matches/{over_match_id}/confirm?operator=tester").status_code == 200

    second_match_response = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": second_bank_id,
            "amount": "600.00",
            "reason": "第二笔付款",
        },
    )
    assert second_match_response.status_code == 409
    assert second_match_response.json()["detail"] == "This bank transaction is already matched to the approval"

    paid_items = client.get("/api/expense-items?payment_status=paid").json()["data"]["items"]
    assert paid_items[0]["id"] == expense_id


def test_expense_match_rejects_income_bank_transaction(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说支出收入方向校验店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "方向校验支出",
            "amount": "100.00",
        },
    ).json()["data"]["id"]
    income_bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "income",
            "amount": "100.00",
            "counterparty_name": "顾客",
        },
    ).json()["data"]["id"]

    response = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": income_bank_id,
            "amount": "100.00",
        },
    )

    assert response.status_code == 409


def test_auto_suggest_match_candidates(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说自动匹配店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "电费",
            "amount": "1280.00",
            "supplier_name": "供电公司",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "1280.00",
            "counterparty_name": "供电公司",
        },
    ).json()["data"]["id"]

    response = client.post("/api/matches/auto-suggest")
    assert response.status_code == 201
    data = response.json()["data"]
    assert data["created_count"] == 1
    assert data["matches"][0]["expense_item_id"] == expense_id
    assert data["matches"][0]["bank_transaction_id"] == bank_id
    assert data["matches"][0]["confidence"] == "98.00"

    duplicate_response = client.post("/api/matches/auto-suggest")
    assert duplicate_response.status_code == 201
    assert duplicate_response.json()["data"]["created_count"] == 0

    logs = client.get("/api/audit-logs?action=match.auto_suggest&page_size=20").json()["data"]["items"]
    assert logs


def test_auto_suggest_matches_partial_paid_remaining_amount(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说部分付款自动匹配店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "分批材料款",
            "amount": "1000.00",
            "supplier_name": "材料供应商",
        },
    ).json()["data"]["id"]
    first_bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "400.00",
            "counterparty_name": "材料供应商",
        },
    ).json()["data"]["id"]
    second_bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-14T10:30:00",
            "direction": "expense",
            "amount": "600.00",
            "counterparty_name": "材料供应商",
        },
    ).json()["data"]["id"]

    first_match_id = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": first_bank_id,
            "amount": "400.00",
            "reason": "第一笔",
        },
    ).json()["data"]["id"]
    client.post(f"/api/matches/{first_match_id}/confirm?operator=tester")

    response = client.post("/api/matches/auto-suggest")
    assert response.status_code == 201
    data = response.json()["data"]
    assert data["created_count"] == 1
    assert data["matches"][0]["expense_item_id"] == expense_id
    assert data["matches"][0]["bank_transaction_id"] == second_bank_id
    assert data["matches"][0]["amount"] == "600.00"


def test_update_bank_transaction_and_audit_log(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说流水维护店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "1280.00",
            "counterparty_name": "供电公司",
            "bank_serial_no": "BANK-UPDATE-001",
        },
    ).json()["data"]["id"]

    update_response = client.patch(
        f"/api/bank-transactions/{bank_id}",
        json={
            "occurred_at": "2026-08-13T11:00:00",
            "amount": "1290.00",
            "summary": "修正电费流水",
            "bank_serial_no": "BANK-UPDATE-002",
        },
    )
    assert update_response.status_code == 200
    data = update_response.json()["data"]
    assert data["amount"] == "1290.00"
    assert data["summary"] == "修正电费流水"
    assert data["bank_serial_no"] == "BANK-UPDATE-002"

    logs = client.get("/api/audit-logs?resource_type=bank_transaction&page_size=20").json()["data"]["items"]
    assert any(log["action"] == "bank_transaction.update" and log["resource_id"] == bank_id for log in logs)


def test_update_bank_transaction_rejects_duplicate_serial_no(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说流水冲突店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    first_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "100.00",
            "bank_serial_no": "BANK-DUP-001",
        },
    ).json()["data"]["id"]
    client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-14T10:30:00",
            "direction": "expense",
            "amount": "200.00",
            "bank_serial_no": "BANK-DUP-002",
        },
    )

    response = client.patch(
        f"/api/bank-transactions/{first_id}",
        json={"bank_serial_no": "BANK-DUP-002"},
    )
    assert response.status_code == 409


def test_update_bank_transaction_rejects_matched_transaction(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说流水匹配店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "广告费",
            "amount": "100.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "100.00",
        },
    ).json()["data"]["id"]
    match_id = client.post(
        "/api/matches",
        json={"expense_item_id": expense_id, "bank_transaction_id": bank_id, "amount": "100.00"},
    ).json()["data"]["id"]
    client.post(f"/api/matches/{match_id}/confirm?operator=tester")

    response = client.patch(f"/api/bank-transactions/{bank_id}", json={"summary": "不应允许修改"})
    assert response.status_code == 409
    assert response.json()["detail"] == "Bank transaction already matched and cannot be edited"


def test_match_requires_same_store_and_period(client: TestClient) -> None:
    store_a = client.post("/api/stores", json={"name": "蘑说 A 店"}).json()["data"]["id"]
    store_b = client.post("/api/stores", json={"name": "蘑说 B 店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_a, "period": "2026-08"})
    client.post("/api/ledgers", json={"store_id": store_b, "period": "2026-08"})

    expense_id = client.post(
        "/api/expense-items",
        json={
            "store_id": store_a,
            "ledger_period": "2026-08",
            "description": "跨店支出",
            "amount": "100.00",
        },
    ).json()["data"]["id"]
    bank_id = client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_b,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "100.00",
        },
    ).json()["data"]["id"]

    response = client.post(
        "/api/matches",
        json={
            "expense_item_id": expense_id,
            "bank_transaction_id": bank_id,
            "amount": "100.00",
        },
    )
    assert response.status_code == 409


def test_ledger_report_summary(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说报表店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "income",
            "amount": "1000.00",
            "counterparty_name": "营业收入",
        },
    )
    client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "物料支出",
            "amount": "260.00",
        },
    )

    response = client.get(f"/api/reports/ledger-summary?store_id={store_id}&period=2026-08")
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["store_name"] == "蘑说报表店"
    assert data["income_amount"] == "1000.00"
    assert data["expense_amount"] == "260.00"
    assert data["profit_amount"] == "740.00"


def test_ledger_report_detail_breakdowns(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说明细报表店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "水电费",
            "amount": "120.00",
            "category_l1": "房租水电",
            "supplier_name": "供电公司",
        },
    )
    client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "广告费",
            "amount": "80.00",
            "category_l1": "营销物料",
        },
    )
    client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "80.00",
            "counterparty_name": "广告公司",
        },
    )

    response = client.get(f"/api/reports/ledger-detail?store_id={store_id}&period=2026-08")
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["summary"]["expense_amount"] == "200.00"
    assert data["category_breakdown"][0]["name"] == "房租水电"
    assert data["supplier_breakdown"][0]["name"] == "供电公司"
    assert len(data["pending_expense_items"]) == 2
    assert len(data["pending_bank_transactions"]) == 1


def test_export_ledger_detail_csv(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说导出报表店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "水电费",
            "amount": "120.00",
            "supplier_name": "供电公司",
        },
    )
    client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "120.00",
            "counterparty_name": "供电公司",
        },
    )

    response = client.get(f"/api/reports/ledger-detail.csv?store_id={store_id}&period=2026-08")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    text = response.text
    assert "账套汇总" in text
    assert "水电费" in text
    assert "供电公司" in text


def test_export_ledger_detail_xlsx(client: TestClient) -> None:
    store_id = client.post("/api/stores", json={"name": "蘑说 Excel 报表店"}).json()["data"]["id"]
    client.post("/api/ledgers", json={"store_id": store_id, "period": "2026-08"})
    client.post("/api/revenue-channels", json={"name": "美团", "requires_bank_match": True})
    client.post(
        "/api/revenue-records",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "revenue_date": "2026-08-01",
            "channel": "美团",
            "gross_amount": "1000.00",
            "net_amount": "980.00",
            "fee_amount": "20.00",
        },
    )
    client.post(
        "/api/expense-items",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "description": "物料采购",
            "amount": "200.00",
            "supplier_name": "食材供应商",
        },
    )
    client.post(
        "/api/bank-transactions",
        json={
            "store_id": store_id,
            "ledger_period": "2026-08",
            "occurred_at": "2026-08-13T10:30:00",
            "direction": "expense",
            "amount": "200.00",
            "counterparty_name": "食材供应商",
        },
    )

    response = client.get(f"/api/reports/ledger-detail.xlsx?store_id={store_id}&period=2026-08")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["账套汇总", "营业收入", "收入渠道汇总", "支出明细", "银行流水"]
    assert workbook["账套汇总"]["A1"].value == "门店"
    assert workbook["账套汇总"]["B1"].value == "蘑说 Excel 报表店"
    assert workbook["营业收入"]["B2"].value == "美团"
    assert workbook["收入渠道汇总"]["A2"].value == "美团"
    assert workbook["收入渠道汇总"]["D2"].value == 20
    expense_descriptions = [
        workbook["支出明细"][f"B{row}"].value
        for row in range(2, workbook["支出明细"].max_row + 1)
    ]
    assert "物料采购" in expense_descriptions
    assert workbook["银行流水"]["D2"].value == "食材供应商"
